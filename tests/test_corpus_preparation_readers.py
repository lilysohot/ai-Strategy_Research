"""I1-2 读取 Adapter 测试：空页/多栏/表格/重复标题/短 MD 反例 + 开发材料 smoke。

验收（任务 I1-2）：fidelity 空页/多栏/表格/重复标题/短 MD 反例通过，不靠装库
声称支持（格式签名 fail-closed、表格冲突显式记账、短 MD 正常处理、重复标题
单元唯一）。不构造真实 PG/模型客户端。
"""

from __future__ import annotations

import io
import json
import subprocess
import sys
from pathlib import Path

import pymupdf
import pytest
from docx import Document

from plugins.corpus.preparation.contract import UnitStatus
from plugins.corpus.preparation.readers import ReaderError, read_document
from plugins.corpus.preparation.readers.pdf_reader import (
    _detect_header_rows,
    _emit_table,
    _garbled_ratio,
    _is_data_like,
    _row_text,
    _Table,
    _TableCell,
    _TableModel,
    _TableRow,
)

REPO = Path(__file__).resolve().parents[1]


# --- Markdown ---


def test_md_short_file_single_paragraph(tmp_path: Path) -> None:
    path = tmp_path / "short.md"
    path.write_text("# 标题\n\n正文一句话。", encoding="utf-8")
    result = read_document(path)
    assert result.format.value == "markdown"
    assert result.page_count is None
    assert [u.kind for u in result.units] == ["heading", "paragraph"]
    assert result.units[1].raw_text == "正文一句话。"
    span = result.units[1].location.char_span
    assert span is not None
    text = path.read_text(encoding="utf-8")
    assert text[span.start : span.end] == "正文一句话。"
    assert result.issues == ()


def test_md_repeated_headings_distinct_units(tmp_path: Path) -> None:
    path = tmp_path / "repeat.md"
    path.write_text("# 风险提示\n\n甲段\n\n# 风险提示\n\n乙段\n", encoding="utf-8")
    result = read_document(path)
    headings = [u for u in result.units if u.kind == "heading"]
    assert len(headings) == 2
    assert headings[0].ordinal != headings[1].ordinal
    assert all(u.status is UnitStatus.KEPT for u in result.units)
    assert not any("ocr" in issue.code for issue in result.issues)


def test_md_structural_kinds_and_offset_fidelity(tmp_path: Path) -> None:
    content = (
        "# 报告\n\n- 第一条\n- 第二条\n\n> 引用内容\n\n"
        "| 列A | 列B |\n|---|---|\n| 1 | 2 |\n\n结尾段落\n"
    )
    path = tmp_path / "structure.md"
    path.write_text(content, encoding="utf-8")
    result = read_document(path)
    kinds = [u.kind for u in result.units]
    assert kinds.count("heading") == 1
    assert kinds.count("list_item") == 2
    assert kinds.count("quote") == 1
    assert kinds.count("table_row") == 3
    assert kinds[-1] == "paragraph"
    text = path.read_text(encoding="utf-8")
    for unit in result.units:
        span = unit.location.char_span
        assert span is not None
        assert text[span.start : span.end] == unit.raw_text, unit.raw_text


def test_md_code_fence_kept_with_reason(tmp_path: Path) -> None:
    path = tmp_path / "code.md"
    path.write_text("```\nkeep line\n```\n", encoding="utf-8")
    result = read_document(path)
    assert result.units[0].raw_text == "```\nkeep line\n```"
    assert result.units[0].reasons == ("code_fence",)


# --- 格式识别 fail-closed（不靠装库声称支持） ---


def test_format_signature_fail_closed(tmp_path: Path) -> None:
    fake_pdf = tmp_path / "fake.pdf"
    fake_pdf.write_text("not a pdf", encoding="utf-8")
    with pytest.raises(ReaderError):
        read_document(fake_pdf)
    fake_docx = tmp_path / "fake.docx"
    fake_docx.write_bytes(b"MZ not a zip container")
    with pytest.raises(ReaderError):
        read_document(fake_docx)
    unsupported = tmp_path / "note.txt"
    unsupported.write_text("x", encoding="utf-8")
    with pytest.raises(ReaderError):
        read_document(unsupported)


# --- DOCX ---


def _png_bytes() -> bytes:
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 4, 4))
    return pixmap.tobytes("png")


def test_docx_repeated_headings_body_order_and_table(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("风险提示", level=1)
    document.add_paragraph("正文甲")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "营收"
    table.cell(1, 1).text = "100"
    document.add_paragraph("表格后段落")
    document.add_heading("风险提示", level=1)
    document.add_paragraph("正文乙")
    path = tmp_path / "synthetic.docx"
    document.save(str(path))

    result = read_document(path)
    kinds = [(u.kind, u.raw_text) for u in result.units]
    assert kinds[0] == ("heading", "风险提示")
    assert kinds[1] == ("paragraph", "正文甲")
    rows = [u for u in result.units if u.kind == "table_row"]
    assert [r.raw_text for r in rows] == ["指标 | 数值", "营收 | 100"]
    assert rows[0].location.cells == ((0, 0), (0, 1))
    # 表格之后的正文顺序保持（不假设 document.paragraphs 覆盖全文）。
    assert kinds[-3:] == [
        ("paragraph", "表格后段落"),
        ("heading", "风险提示"),
        ("paragraph", "正文乙"),
    ]
    headings = [u for u in result.units if u.kind == "heading"]
    assert headings[0].ordinal != headings[1].ordinal


def test_docx_image_records_gap_issue(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("图片前段落")
    document.add_picture(io.BytesIO(_png_bytes()))
    document.add_paragraph("图片后段落")
    path = tmp_path / "with-image.docx"
    document.save(str(path))

    result = read_document(path)
    assert "unreadable_element" in [issue.code for issue in result.issues]
    texts = [u.raw_text for u in result.units]
    assert "图片前段落" in texts and "图片后段落" in texts


# --- PDF ---


def test_pdf_paragraphs_headings_bbox(tmp_path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Overview", fontsize=16)
    page.insert_text((72, 140), "First body line", fontsize=11)
    page.insert_text((72, 155), "Second body line", fontsize=11)
    page.insert_text((72, 220), "Separate paragraph", fontsize=11)
    path = tmp_path / "text.pdf"
    doc.save(str(path))
    doc.close()

    result = read_document(path)
    assert result.page_count == 1
    headings = [u for u in result.units if u.kind == "heading"]
    assert [h.raw_text for h in headings] == ["Overview"]
    assert headings[0].location.page == 1
    assert headings[0].location.bbox is not None
    paragraphs = [u for u in result.units if u.kind == "paragraph"]
    assert len(paragraphs) == 2
    assert paragraphs[0].raw_text == "First body line\nSecond body line"
    assert paragraphs[1].raw_text == "Separate paragraph"


def test_pdf_empty_and_image_only_pages(tmp_path: Path) -> None:
    doc = pymupdf.open()
    doc.new_page()  # 第 1 页：无文字层无图片
    doc.new_page().insert_text((72, 72), "hello", fontsize=11)
    image_page = doc.new_page()
    image_page.insert_image(
        pymupdf.Rect(100, 100, 200, 200),
        pixmap=pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 10, 10)),
    )
    path = tmp_path / "pages.pdf"
    doc.save(str(path))
    doc.close()

    result = read_document(path)
    assert result.page_count == 3
    codes = [issue.code for issue in result.issues]
    assert codes.count("empty_page") == 1
    assert codes.count("image_only_page") == 1
    assert result.units and all(u.location.page == 2 for u in result.units)


def test_pdf_multicolumn_flagged_and_covered(tmp_path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    for index, y in enumerate((80, 100, 120, 140)):
        page.insert_text((60, y), f"LEFT{index} alpha beta", fontsize=11)
        page.insert_text((360, y), f"RIGHT{index} gamma delta", fontsize=11)
    path = tmp_path / "columns.pdf"
    doc.save(str(path))
    doc.close()

    result = read_document(path)
    assert "multi_column_order_unreliable" in [issue.code for issue in result.issues]
    order = [u.raw_text for u in result.units]
    joined = "\n".join(order)
    assert "LEFT0" in joined and "RIGHT3" in joined
    left_first = next(i for i, text in enumerate(order) if "LEFT0" in text)
    right_first = next(i for i, text in enumerate(order) if "RIGHT0" in text)
    assert left_first < right_first  # 阅读次序：先左栏后右栏


def test_pdf_true_two_column_native_order_kept(tmp_path: Path) -> None:
    """真双栏且内容流按整栏写入：原生序本身即「先左栏后右栏」，不需要强制重排。

    这是关闭 `_split_columns` 强制左→右重排后的正对照：证明移除重排不会让真正的
    双栏文档串行（I3-3 关闭重排的唯一语义风险）。
    """
    doc = pymupdf.open()
    page = doc.new_page()
    ys = (80, 100, 120, 140)
    for index, y in enumerate(ys):  # 整栏写：先写满左栏，再写满右栏
        page.insert_text((60, y), f"LEFT{index} alpha beta", fontsize=11)
    for index, y in enumerate(ys):
        page.insert_text((360, y), f"RIGHT{index} gamma delta", fontsize=11)
    path = tmp_path / "columns_native.pdf"
    doc.save(str(path))
    doc.close()

    result = read_document(path)
    assert "multi_column_order_unreliable" in [issue.code for issue in result.issues]
    order = [u.raw_text for u in result.units]
    left_idx = [i for i, text in enumerate(order) if "LEFT" in text]
    right_idx = [i for i, text in enumerate(order) if "RIGHT" in text]
    assert left_idx and right_idx
    assert max(left_idx) < min(right_idx)  # 原生序已是先左栏后右栏


def test_pdf_table_rows_no_silent_duplication(tmp_path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    xs = [50, 150, 250, 350]
    ys = [300, 330, 360, 390]
    for x in xs:
        page.draw_line((x, ys[0]), (x, ys[-1]))
    for y in ys:
        page.draw_line((xs[0], y), (xs[-1], y))
    for column, header in enumerate(["Yr", "Rev", "NP"]):
        page.insert_text(((xs[column] + xs[column + 1]) / 2 - 6, ys[0] + 20), header, fontsize=9)
    for row_index, row in enumerate([["2024", "9901", "11"], ["2025", "9902", "12"]]):
        for column, value in enumerate(row):
            page.insert_text(
                ((xs[column] + xs[column + 1]) / 2 - 6, ys[row_index + 1] + 20), value, fontsize=9
            )
    path = tmp_path / "table.pdf"
    doc.save(str(path))
    doc.close()

    result = read_document(path)
    rows = [u for u in result.units if u.kind == "table_row"]
    assert len(rows) >= 3
    table_text = "\n".join(u.raw_text for u in rows)
    assert "2024" in table_text and "9901" in table_text and "9902" in table_text
    assert "|" not in table_text  # I1：不得插入原文没有的字符（旧实现插入 " | "）
    prose = [u for u in result.units if u.kind != "table_row"]
    assert not any("9901" in u.raw_text or "9902" in u.raw_text for u in prose)
    assert "table_text_overlap" in [issue.code for issue in result.issues]


def test_table_row_text_invariants_native_order_and_source_characters() -> None:
    """表格行发射的两条不变量（I3-3 表格类引文失败的共因）。

    I1（字符保真）：单元格之间只用换行，不插入原文没有的字符。
    I2（顺序保真）：文本顺序服从原生阅读序；网格行列只作为 locator 元数据下发，
    不参与文本重排。
    """
    row = _TableRow(
        cells=(
            _TableCell(text="产品", row=0, col=1, native_pos=2),
            _TableCell(text="产能（万吨/年）", row=0, col=9, native_pos=1),  # 内容流里先画
            _TableCell(text="价格分位", row=0, col=6, native_pos=3),
        )
    )
    text, cells = _row_text(row)
    assert text == "产能（万吨/年）\n产品\n价格分位"  # 原生序，而非网格列序
    assert "|" not in text
    assert cells == ((0, 9), (0, 1), (0, 6))  # 行列仍在，只是降级为元数据


def test_garbled_ratio_helper() -> None:
    assert _garbled_ratio("") == 0.0
    assert _garbled_ratio("正常文本") == 0.0
    assert _garbled_ratio("abc\ufffd\ufffd") == pytest.approx(0.4)


# --- 表格结构模型（票 04：确定性表头识别 + 合并单元格 + label_path） ---


def test_is_data_like_discriminates_values_from_period_years() -> None:
    assert _is_data_like("5814.2")
    assert _is_data_like("32.1")
    assert _is_data_like("89.9%")
    assert _is_data_like("-1,234.5")
    assert not _is_data_like("尿素")
    assert not _is_data_like("2023 2024 2025")  # 裸 4 位年份是期间标签，非数据值


def test_detect_header_rows_stops_at_data_rows_and_skips_sparse() -> None:
    # 数据形态占比 ≥50% 即停；稀疏行（整行合并标题）跳过不计数
    grid = [
        ["开工率", "价差", "产能", "价格分位"],
        ["尿素", "5814.2", "32.1", "89.9%"],
        ["纯碱", "6728.3", "10.9", "82.9%"],
    ]
    assert _detect_header_rows(grid) == (0,)
    sparse = [
        ["表 6 主要化工品景气跟踪"],  # 稀疏行：非空 <2，跳过
        ["产品", "开工率", "产能"],
        ["尿素", "32.1", "100"],
    ]
    assert _detect_header_rows(sparse) == (1,)
    # 表头期间行（裸 4 位年份）豁免为表头而非数据行
    with_years = [
        ["产品", "2023", "2024", "2025"],
        ["尿素", "5814.2", "32.1", "89.9%"],
    ]
    assert _detect_header_rows(with_years) == (0,)


def test_table_model_label_path_merged_headers() -> None:
    """I-B1 读侧投影：多级表头 + 合并单元格按 bbox 覆盖继承列标签。

    ``开工率`` 跨列 2-3 合并（覆盖位 None）；``col_labels`` 自下而上遍历表头行，
    数据行锚定单元格的列中心落入合并 bbox 即继承标签。
    """
    model = _TableModel(
        page=10,
        table_index=0,
        header_rows=(0, 1),
        grid=(
            (None, "产品", "开工率", None),  # 合并单元格：开工率 横跨列 2-3
            ("周期", "尿素", "纯碱", "纯碱"),
            (None, "尿素", "5814.2", "8068.0"),  # 数据行
        ),
        anchored=(
            ((1, (100.0, 200.0)), (2, (200.0, 400.0))),  # 合并 bbox 覆盖列 2-3
            ((0, (50.0, 100.0)), (1, (100.0, 200.0)), (2, (200.0, 300.0)), (3, (300.0, 400.0))),
            ((1, (100.0, 200.0)), (2, (200.0, 300.0)), (3, (300.0, 400.0))),
        ),
    )
    # 数据行 cell(2, 3)：行标签 尿素 + 列标签链（纯碱 ← 开工率）
    assert model.label_path(2, 3) == ("尿素", "纯碱", "开工率")
    # cell(2, 2) 同属纯碱 列，且开工率 合并 bbox 覆盖列 2-3 → 同样继承完整链
    assert model.label_path(2, 2) == ("尿素", "纯碱", "开工率")
    # 表头行无行标签
    assert model.label_path(1, 2) == ("纯碱", "开工率")
    # 覆盖位空单元格不猜测内容
    assert model.cell_text(0, 3) == ""


def test_emit_table_model_label_path_aligned_and_row_text_untouched() -> None:
    """保真反例：带 model 的表格发射不触碰原文。

    ``raw_text`` 与 ``_row_text`` 逐字节一致（仍按 native_pos 排序、只用换行连接，
    不插入 ``" | "``）；``label_path`` 与 ``cells`` 一一对齐（I-B1 读侧落点）。
    """
    model = _TableModel(
        page=10,
        table_index=0,
        header_rows=(0,),
        grid=(
            ("产品", "开工率", "产能"),
            ("尿素", "89.9%", "100"),
        ),
        anchored=(
            ((0, (50.0, 100.0)), (1, (100.0, 200.0)), (2, (200.0, 300.0))),
            ((0, (50.0, 100.0)), (1, (100.0, 200.0)), (2, (200.0, 300.0))),
        ),
    )
    row = _TableRow(
        cells=(
            _TableCell(text="产能", row=1, col=2, native_pos=0),
            _TableCell(text="尿素", row=1, col=0, native_pos=1),
            _TableCell(text="89.9%", row=1, col=1, native_pos=2),
        )
    )
    table = _Table(index=0, bbox=(50.0, 300.0, 300.0, 330.0), rows=[row], model=model)
    units: list[object] = []
    issues: list[object] = []
    _emit_table(table, 0, page_no=10, page_lines=[], units=units, issues=issues)
    assert len(units) == 1
    unit = units[0]  # type: ignore[attr-defined]
    text, cells = _row_text(row)
    assert unit.raw_text == text == "产能\n尿素\n89.9%"  # 原生序 + 纯换行，不插 " | "
    assert unit.location.cells == cells == ((1, 2), (1, 0), (1, 1))
    assert len(unit.location.label_path) == len(cells)
    # label_path 与 cells 顺序对齐（原生序）：产能→(尿素, 产能)、尿素→(尿素, 产品)、89.9%→(尿素, 开工率)
    assert unit.location.label_path == ("尿素 产能", "尿素 产品", "尿素 开工率")
    assert "|" not in unit.raw_text


# --- 开发材料 smoke（守卫允许清单内的 6 份真实 PDF，只读） ---


def _dev_materials() -> list[Path]:
    config_path = REPO / ".scratch/corpus-evidence-pipeline/ingestion-rebuild/guards/i1.json"
    config = json.loads(config_path.read_text(encoding="utf-8"))
    return [REPO / raw for raw in config["sources"]["allowed_source_paths"]]


def test_dev_materials_smoke_units_and_pages() -> None:
    for path in _dev_materials():
        result = read_document(path)
        assert result.page_count is not None and result.page_count >= 1, path.name
        assert result.units, path.name
        assert result.extractor_rev.startswith("reader-pdf-6+")


def test_readers_do_not_import_pg_or_model_client() -> None:
    # 干净子进程断言：同进程其他测试文件不应污染本纪律
    code = (
        "import sys; import plugins.corpus.preparation.readers;"
        "bad = [m for m in ('psycopg', 'psycopg2', 'asyncpg', 'sqlalchemy',"
        " 'openai', 'anthropic') if m in sys.modules];"
        "print(f'leaked: {bad}'); sys.exit(1 if bad else 0)"
    )
    proc = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, check=False)
    assert proc.returncode == 0, proc.stdout + proc.stderr
