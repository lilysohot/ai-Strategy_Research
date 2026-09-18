"""I1 复核修复回归（remediation：20260916-i1 review 的 F1—F7）。

15 个复核反例转正为正式回归测试（验收探针原件见
``.scratch/corpus-evidence-pipeline/ingestion-rebuild/audits/20260916-i1/``）：
断言与探针一致，不得调整断言适配已知错误；历史通过记录不作废。全部为合成
输入，不触真实语料，不构造真实 PG/模型客户端。
"""

from __future__ import annotations

from dataclasses import replace
from datetime import UTC, datetime, timedelta

import pymupdf
import pytest
from docx import Document

from plugins.corpus.preparation.admission import (
    POLICY_REV_V1,
    AdmissionPolicy,
    ProbeInput,
    SourceDescriptor,
    decide_admission,
    probe_features,
)
from plugins.corpus.preparation.chunk import chunk_clean_result
from plugins.corpus.preparation.clean import CleanError, clean_reader_result, verify_clean_region
from plugins.corpus.preparation.contract import (
    Admission,
    AdmissionDecision,
    Build,
    DocumentFormat,
    JobStage,
    LeaseConfig,
    MaterialType,
    ResearchDomain,
    ReviewDecision,
    ReviewedDecision,
    Unit,
    compute_build_id,
    source_id_from_bytes,
)
from plugins.corpus.preparation.readers import read_document
from plugins.corpus.preparation.repository import MemoryStore, StoreError

NOW = datetime(2026, 9, 16, tzinfo=UTC)
SID = "a" * 64


def _setup_store() -> tuple[MemoryStore, Admission, Build]:
    store = MemoryStore()
    admission = Admission(
        "d0",
        SID,
        MaterialType.RESEARCH_REPORT,
        ResearchDomain.COMPANY,
        AdmissionDecision.IN_SCOPE,
    )
    args = dict(
        source_id=SID,
        decision_id="d0",
        parse_rev="p1",
        clean_rev="c1",
        chunk_rev="k1",
        index_rev="i1",
        scope_ref=None,
    )
    build = Build(build_id=compute_build_id(**args), **args)
    store.put_admission(admission)
    store.put_build(build)
    return store, admission, build


def _decide_with(reviews: tuple[ReviewedDecision, ...]) -> Admission:
    source = SourceDescriptor(
        SID, DocumentFormat.MARKDOWN, "synthetic.md", "证券公司研究", ResearchDomain.COMPANY
    )
    probes = probe_features(ProbeInput(source.title, "证券研究所", ("正文。",), ()))
    policy = AdmissionPolicy(POLICY_REV_V1, tuple(ResearchDomain))
    return decide_admission(source, reviews, probes, policy)


def _review(
    name: str, decision: ReviewDecision = ReviewDecision.ADMITTED, supersedes: str | None = None
) -> ReviewedDecision:
    return ReviewedDecision(
        name, SID, "reviewer", NOW, decision, "synthetic approval", supersedes=supersedes
    )


# --- F1：put_admission 历史追加与当前指针分离；publish 核对当前决定/scope ---


def test_old_admission_replay_must_not_rollback_current() -> None:
    store, old, _ = _setup_store()
    revoked = replace(old, decision_id="d1", decision=AdmissionDecision.EXCLUDED_BY_POLICY)
    store.put_admission(revoked)
    store.put_admission(old)  # 幂等重放历史记录不得回退当前指针
    assert store.latest_admission(SID).decision_id == "d1"


def _acquire_published_lease(store: MemoryStore, build: Build) -> str:
    """与业务走同一发布协议：登记 + 取得 PUBLISHED job 当前租约，返回 fence_token。

    acquire 用真实墙钟（MemoryStore 默认时钟判定租约过期，固定 NOW 会被判过期）。
    """
    lease = LeaseConfig(120, 30, 360, 3)
    store.register_job(build.build_id, JobStage.PUBLISHED)
    acquired = store.acquire_job(
        build.build_id, JobStage.PUBLISHED, "fixture", datetime.now(UTC), lease
    )
    token = acquired.fence_token
    assert token is not None
    return token


def test_revoked_source_cannot_be_republished_by_old_decision() -> None:
    store, old, build = _setup_store()
    store.put_admission(
        replace(old, decision_id="d1", decision=AdmissionDecision.EXCLUDED_BY_POLICY)
    )
    store.retire(SID, "d1", NOW)
    token = _acquire_published_lease(store, build)
    with pytest.raises(StoreError):
        store.publish(SID, "d0", build.build_id, NOW, owner_id="fixture", fence_token=token)


def test_scoped_approval_cannot_publish_full_scope_build() -> None:
    store, old, build = _setup_store()
    store.put_admission(replace(old, decision_id="d1", scope_ref="page:1"))
    token = _acquire_published_lease(store, build)
    with pytest.raises(StoreError):
        store.publish(SID, "d1", build.build_id, NOW, owner_id="fixture", fence_token=token)


# --- F2：租约接管尊重 max_attempts；权威输出写入绑定当前租约 ---


def test_expired_running_jobs_respect_max_attempts() -> None:
    store, _, build = _setup_store()
    lease = LeaseConfig(
        lease_ttl_seconds=120,
        heartbeat_interval_seconds=30,
        stage_timeout_seconds=360,
        max_attempts=2,
    )
    store.register_job(build.build_id, JobStage.PARSED)
    store.acquire_job(build.build_id, JobStage.PARSED, "a", NOW, lease)
    store.acquire_job(build.build_id, JobStage.PARSED, "b", NOW + timedelta(seconds=121), lease)
    with pytest.raises(StoreError):
        store.acquire_job(build.build_id, JobStage.PARSED, "c", NOW + timedelta(seconds=242), lease)


def test_authoritative_output_cannot_be_written_without_current_ownership() -> None:
    store, _, build = _setup_store()
    lease = LeaseConfig(
        lease_ttl_seconds=120,
        heartbeat_interval_seconds=30,
        stage_timeout_seconds=360,
        max_attempts=2,
    )
    store.register_job(build.build_id, JobStage.PARSED)
    store.acquire_job(build.build_id, JobStage.PARSED, "a", NOW, lease)
    store.acquire_job(build.build_id, JobStage.PARSED, "b", NOW + timedelta(seconds=121), lease)
    unit = Unit(
        "late-a", build.build_id, "paragraph", "late output", source_id_from_bytes(b"late output")
    )
    with pytest.raises(StoreError):
        store.put_units(build.build_id, [unit])  # 无当前租约 owner/token 不得直写权威输出


# --- F5：取代链全连通分量检查；非法决定词表构造即拒绝 ---


def test_disconnected_cycle_is_not_a_resolved_review_history() -> None:
    result = _decide_with(
        (_review("a", supersedes="b"), _review("b", supersedes="a"), _review("c"))
    )
    assert result.decision is AdmissionDecision.REVIEW_REQUIRED


def test_unknown_review_enum_cannot_be_treated_as_approval() -> None:
    with pytest.raises(ValueError):
        _review("pending-review", decision="pending")  # type: ignore[arg-value]


# --- F3：PDF 同页标题/正文交错次序保持 ---


def test_pdf_units_preserve_interleaved_heading_body_order(tmp_path) -> None:
    path = tmp_path / "sections.pdf"
    with pymupdf.open() as doc:
        page = doc.new_page()
        for y, text, size in [
            (72, "Section A", 18),
            (110, "A body 1", 11),
            (145, "A body 2", 11),
            (215, "Section B", 18),
            (250, "B body 1", 11),
            (285, "B body 2", 11),
        ]:
            page.insert_text((72, y), text, fontsize=size)
        doc.save(str(path))
    result = read_document(path)
    actual = [u.raw_text for u in result.units]
    assert actual.index("A body 1") < actual.index("Section B"), actual


# --- F4：混合图片页与 DOCX 嵌套表格不静默丢失 ---


def test_pdf_image_region_is_accounted_for_even_with_text(tmp_path) -> None:
    path = tmp_path / "mixed-page.pdf"
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_text((72, 72), "Report title", fontsize=11)
        page.insert_image(
            pymupdf.Rect(72, 120, 520, 700),
            pixmap=pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 10, 10)),
        )
        doc.save(str(path))
    result = read_document(path)
    assert result.issues or any(u.status.value != "kept" for u in result.units), result


def test_docx_nested_table_is_extracted_or_recorded_as_gap(tmp_path) -> None:
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Outer cell"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Revenue 1234"
    path = tmp_path / "nested.docx"
    document.save(str(path))
    result = read_document(path)
    assert "Revenue 1234" in "\n".join(u.raw_text for u in result.units) or result.issues, result


# --- F6：保真校验拒绝擦除与重排 ---


@pytest.mark.parametrize(
    "clean,mapping", [("", ()), ("21", ((0, 1), (1, 0)))], ids=["erasure", "reordering"]
)
def test_clean_verifier_rejects_erasure_and_reordering(clean, mapping) -> None:
    with pytest.raises(CleanError):
        verify_clean_region("12", clean, mapping)


# --- F7：MD 表序分组、末尾标题保留、context 不挤爆块上限 ---


def test_heading_only_document_is_preserved(tmp_path) -> None:
    path = tmp_path / "heading.md"
    path.write_text("# Heading only\n", encoding="utf-8")
    reader = read_document(path)
    result = chunk_clean_result(reader, clean_reader_result(reader))
    assert len(result.chunks) == 1
    assert "Heading only" in result.chunks[0].search_text


def test_adjacent_markdown_tables_are_not_fused(tmp_path) -> None:
    path = tmp_path / "tables.md"
    path.write_text(
        "| 2025 | revenue |\n|---|---|\n| A | 100 |\n\n| 2026 | margin |\n|---|---|\n| B | 20% |\n",
        encoding="utf-8",
    )
    reader = read_document(path)
    result = chunk_clean_result(reader, clean_reader_result(reader))
    tables = [chunk for chunk in result.chunks if chunk.kind == "table"]
    assert len(tables) == 2, tables


def test_table_header_context_does_not_make_valid_rows_unprocessable(tmp_path) -> None:
    path = tmp_path / "long-rows.md"
    path.write_text("| " + "H" * 990 + " |\n|---|\n| " + "A" * 990 + " |\n", encoding="utf-8")
    reader = read_document(path)
    result = chunk_clean_result(reader, clean_reader_result(reader))
    assert result.chunks
    assert all(len(c.search_text) <= 1800 or c.review_reasons for c in result.chunks)
