"""R2 material semantics: attribution, dialogue, relations and fail-closed evidence."""

from __future__ import annotations

import json
from pathlib import Path
from unittest.mock import Mock

from plugins.corpus.claims_v2 import triage_block_detail
from plugins.corpus.evidence import split_spans
from plugins.corpus.evidence_pipeline import build_evidence_run
from plugins.corpus.material_semantics import (
    MaterialEvidence,
    MaterialItem,
    _relation_candidate_pairs,
    _relations_from_decisions,
    build_candidate_slot_batches,
    build_candidate_slots,
    build_material_structure,
    classify_material_type,
    extract_material_understanding,
)
from plugins.corpus.service import CorpusService


def _source(tmp_path: Path) -> Path:
    source = tmp_path / "5月：锂电铜箔和电子铜箔.md"
    source.write_text(
        "# 锂电铜箔和电子铜箔\n\n"
        "主持人：如果客户验证顺利，后续需求会明显起来吗？\n"
        "专家：如果四季度验证顺利，我判断明年需求会明显起来。\n"
        "投资者：这个判断有什么依据？\n"
        "专家：铜冠负责人说过，国内设备是六十四，海外是四十。\n",
        encoding="utf-8",
    )
    return source


def _response(_prompt: str) -> str:
    return json.dumps(
        {
            "speakers": [
                {
                    "speaker_id": "host",
                    "display_name": "主持人",
                    "role": "host",
                    "identity_status": "unknown",
                },
                {
                    "speaker_id": "expert",
                    "display_name": "专家",
                    "role": "expert",
                    "identity_status": "unknown",
                },
                {
                    "speaker_id": "investor",
                    "display_name": "投资者",
                    "role": "questioner",
                    "identity_status": "unknown",
                },
                {
                    "speaker_id": "quoted",
                    "display_name": "铜冠负责人",
                    "role": "quoted_source",
                    "identity_status": "unknown",
                },
            ],
            "items": [
                {
                    "item_id": "q1",
                    "text": "主持人询问客户验证顺利后需求是否会起来",
                    "semantic_type": "unknown",
                    "statement_role": "question",
                    "speech_role": "question",
                    "perspective": "source_explicit",
                    "speaker_ref": "host",
                    "polarity": "unknown",
                    "value": None,
                    "behavior_status": None,
                    "temporal_frame": "contemporaneous",
                    "evidence_quote": "主持人：如果客户验证顺利，后续需求会明显起来吗？",
                    "unknown_fields": [],
                },
                {
                    "item_id": "a1",
                    "text": "若四季度验证顺利，专家判断明年需求会明显起来",
                    "semantic_type": "forecast",
                    "statement_role": "condition",
                    "speech_role": "answer",
                    "perspective": "source_explicit",
                    "speaker_ref": "expert",
                    "polarity": "affirmed",
                    "value": None,
                    "behavior_status": None,
                    "temporal_frame": "contemporaneous",
                    "evidence_quote": "专家：如果四季度验证顺利，我判断明年需求会明显起来。",
                    "unknown_fields": [],
                },
                {
                    "item_id": "quoted1",
                    "text": "铜冠负责人被转述称国内设备是六十四、海外是四十",
                    "semantic_type": "fact",
                    "statement_role": "evidence",
                    "speech_role": "answer",
                    "perspective": "quoted_other",
                    "speaker_ref": "quoted",
                    "polarity": "affirmed",
                    "value": "国内六十四；海外四十",
                    "behavior_status": None,
                    "temporal_frame": "retrospective",
                    "evidence_quote": "铜冠负责人说过，国内设备是六十四，海外是四十。",
                    "unknown_fields": ["unit", "observation_date"],
                },
            ],
            "relations": [
                {
                    "relation_id": "r1",
                    "type": "answers",
                    "from_item": "a1",
                    "to_item": "q1",
                    "provenance": "source_explicit",
                    "evidence_quote": "主持人：如果客户验证顺利，后续需求会明显起来吗？\n"
                    "专家：如果四季度验证顺利，我判断明年需求会明显起来。",
                }
            ],
        },
        ensure_ascii=False,
    )


def _item_jsonl() -> str:
    payload = json.loads(_response(""))
    records = [
        *({"record_type": "speaker", **speaker} for speaker in payload["speakers"]),
        *({"record_type": "item", **item} for item in payload["items"]),
    ]
    return "\n".join(json.dumps(record, ensure_ascii=False) for record in records)


def test_dialogue_without_digits_is_a_candidate() -> None:
    decision = triage_block_detail("主持人：怎么看需求？\n专家：我认为需求会改善。")
    assert decision.candidate and decision.reason == "dialogue"


def test_material_semantics_preserve_attribution_conditions_and_quotes(tmp_path: Path) -> None:
    evidence_run = build_evidence_run(_source(tmp_path), packet_chars=1000)
    material_run = extract_material_understanding(evidence_run, llm=_response, max_calls=1)
    material_run.verify_identity()

    result = material_run.understanding
    assert result.source.material_type == "conference_minutes"
    assert result.source.research_domain == "industry"
    assert result.source.material_date is None
    assert len(result.speakers) == 5 and all(
        s.identity_status == "unknown" for s in result.speakers
    )
    assert next(s for s in result.speakers if s.display_name == "专家").role == "industry_expert"
    assert [item.semantic_type for item in result.items] == ["unknown", "forecast", "fact"]
    assert result.items[1].statement_role == "condition"
    assert result.items[1].speech_role == "answer"
    assert result.items[1].value is None
    assert "numeric_value" not in result.items[1].unknown_fields
    assert result.items[2].perspective == "quoted_other"
    assert (
        next(
            speaker
            for speaker in result.speakers
            if speaker.speaker_id == result.items[2].speaker_ref
        ).role
        == "quoted_source"
    )
    assert result.relations[0].provenance == "source_explicit"
    assert not result.coverage.omitted_areas
    assert material_run.summary()["complete"] is True
    for item in result.items:
        assert item.evidence[0].source_rev == evidence_run.document.source_rev
        assert item.evidence[0].quote in evidence_run.document.packets[0].text


def test_bad_speaker_reference_discards_only_dependent_records(tmp_path: Path) -> None:
    payload = json.loads(_response(""))
    payload["items"][0]["speaker_ref"] = "invented"
    evidence_run = build_evidence_run(_source(tmp_path), packet_chars=1000)
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: json.dumps(payload, ensure_ascii=False),
        max_calls=1,
    )
    assert len(result.understanding.items) == 2
    assert all("主持人询问" not in item.text for item in result.understanding.items)
    assert not result.understanding.relations
    assert result.packet_runs[0].status == "completed"
    assert result.packet_runs[0].diagnostics == {"discarded_records": 2}


def test_nonbehavior_status_is_normalized_to_not_applicable(tmp_path: Path) -> None:
    payload = json.loads(_response(""))
    payload["items"][1]["behavior_status"] = "not_applicable"
    evidence_run = build_evidence_run(_source(tmp_path), packet_chars=1000)
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: json.dumps(payload, ensure_ascii=False),
        max_calls=1,
    )
    assert result.packet_runs[0].status == "completed"
    assert result.understanding.items[1].behavior_status is None


def test_ambiguous_extra_quote_does_not_discard_valid_packet_items(tmp_path: Path) -> None:
    payload = json.loads(_response(""))
    payload["items"].append({**payload["items"][0], "item_id": "bad", "evidence_quote": "专家"})
    evidence_run = build_evidence_run(_source(tmp_path), packet_chars=1000)
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: json.dumps(payload, ensure_ascii=False),
        max_calls=1,
    )
    assert len(result.understanding.items) == 3
    assert result.packet_runs[0].status == "completed"
    assert result.packet_runs[0].diagnostics == {"discarded_records": 1}


def test_invalid_extra_enum_does_not_discard_valid_packet_items(tmp_path: Path) -> None:
    payload = json.loads(_response(""))
    payload["items"].append({**payload["items"][0], "item_id": "bad", "semantic_type": "condition"})
    evidence_run = build_evidence_run(_source(tmp_path), packet_chars=1000)
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: json.dumps(payload, ensure_ascii=False),
        max_calls=1,
    )
    assert len(result.understanding.items) == 3
    assert result.packet_runs[0].status == "completed"
    assert result.packet_runs[0].diagnostics == {"discarded_records": 1}


def test_model_enum_aliases_are_normalized_at_the_contract_boundary(tmp_path: Path) -> None:
    payload = json.loads(_response(""))
    payload["items"][0].update(
        semantic_type="risk",
        statement_role="risk",
        polarity="negative",
        temporal_frame="future",
    )
    payload["items"][1].update(polarity="positive", temporal_frame="2026Q4")
    payload["items"][2].update(
        semantic_type="behavior",
        behavior_status="completed",
        temporal_frame="yesterday",
    )
    evidence_run = build_evidence_run(_source(tmp_path), packet_chars=1000)
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: json.dumps(payload, ensure_ascii=False),
        max_calls=1,
    )
    first, second, third = result.understanding.items
    assert first.semantic_type == "forecast"
    assert first.polarity == "unknown"
    assert first.temporal_frame == "unknown"
    assert second.polarity == "affirmed"
    assert second.temporal_frame == "unknown"
    assert third.behavior_status == "claimed_executed"
    assert third.temporal_frame == "retrospective"


def test_truncated_response_salvages_complete_records_and_marks_partial(tmp_path: Path) -> None:
    raw = _response("")
    truncated = raw[: raw.index('"relations"')] + '"relations": [{"relation_id": "cut"'
    evidence_run = build_evidence_run(_source(tmp_path), packet_chars=1000)
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: truncated,
        max_calls=1,
    )
    assert len(result.understanding.items) == 3
    assert result.packet_runs[0].status == "partial"
    assert result.packet_runs[0].diagnostics == {"partial_json_salvaged": True}
    assert result.understanding.coverage.omitted_areas
    assert result.summary()["complete"] is False


def test_structured_split_keeps_question_with_following_answer() -> None:
    text = (
        "专家：" + "前文" * 30 + "。\n\n"
        "主持人：后续需求怎么看？\n\n"
        "专家：如果验证成功，需求会起来。\n\n"
        "主持人：第二个问题？\n\n"
        "专家：第二个回答。"
    )
    spans = split_spans(text, "document", 120)
    assert "".join(span.text for span in spans) == text
    question_packet = next(span.text for span in spans if "后续需求怎么看" in span.text)
    assert "如果验证成功，需求会起来" in question_packet


def test_company_announcement_and_target_price_are_normalized(tmp_path: Path) -> None:
    source = tmp_path / "company.md"
    source.write_text(
        "事项：公司公布中报，26H1实现收入100亿元。\n投资建议：维持目标价2030元。\n",
        encoding="utf-8",
    )
    payload = {
        "speakers": [
            {
                "speaker_id": "analyst",
                "display_name": "分析师甲",
                "role": "analyst_author",
                "identity_status": "explicit",
            }
        ],
        "items": [
            {
                "item_id": "fact",
                "text": "26H1实现收入100亿元",
                "semantic_type": "fact",
                "statement_role": "claim",
                "speech_role": "statement",
                "perspective": "source_explicit",
                "speaker_ref": "analyst",
                "polarity": "affirmed",
                "value": "100亿元",
                "behavior_status": None,
                "temporal_frame": "retrospective",
                "evidence_quote": "26H1实现收入100亿元",
                "unknown_fields": [],
            },
            {
                "item_id": "target",
                "text": "维持目标价2030元",
                "semantic_type": "opinion",
                "statement_role": "claim",
                "speech_role": "statement",
                "perspective": "source_explicit",
                "speaker_ref": "analyst",
                "polarity": "affirmed",
                "value": "2030元",
                "behavior_status": None,
                "temporal_frame": "contemporaneous",
                "evidence_quote": "维持目标价2030元",
                "unknown_fields": [],
            },
        ],
        "relations": [],
    }
    evidence_run = build_evidence_run(source, packet_chars=1000)
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: json.dumps(payload, ensure_ascii=False),
        max_calls=1,
    ).understanding
    fact, target = result.items
    fact_speaker = next(s for s in result.speakers if s.speaker_id == fact.speaker_ref)
    assert fact.statement_role == "evidence"
    assert fact.perspective == "quoted_other"
    assert fact_speaker.role == "quoted_source"
    assert target.semantic_type == "forecast"


def test_staged_jsonl_extracts_items_then_relations(tmp_path: Path) -> None:
    calls: list[str] = []

    def llm(prompt: str) -> str:
        calls.append(prompt)
        if "本阶段只抽 speakers 和 items" in prompt:
            return _item_jsonl()
        marker = "可用 items：\n"
        catalog = json.loads(prompt.split(marker, 1)[1].split("\n\n当前证据包：", 1)[0])
        question = next(item for item in catalog if item["speech_role"] == "question")
        answer = next(item for item in catalog if item["speech_role"] == "answer")
        return json.dumps(
            {
                "record_type": "relation",
                "relation_id": "r1",
                "type": "answers",
                "from_item": answer["item_id"],
                "to_item": question["item_id"],
                "provenance": "source_explicit",
                "evidence_quote": "主持人：如果客户验证顺利，后续需求会明显起来吗？\n"
                "专家：如果四季度验证顺利，我判断明年需求会明显起来。",
            },
            ensure_ascii=False,
        )

    evidence_run = build_evidence_run(_source(tmp_path), packet_chars=1000)
    result = extract_material_understanding(
        evidence_run,
        llm=llm,
        max_calls=2,
        staged_jsonl=True,
        max_items_per_packet=30,
    )
    assert len(calls) == 2
    assert len(result.understanding.items) == 3
    assert len(result.understanding.relations) == 1
    assert result.packet_runs[0].status == "completed"
    assert result.packet_runs[0].model_calls == 2


def test_staged_jsonl_salvages_complete_lines_and_marks_partial(tmp_path: Path) -> None:
    responses = iter((_item_jsonl() + '\n{"record_type":"item"', ""))
    evidence_run = build_evidence_run(_source(tmp_path), packet_chars=1000)
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: next(responses),
        max_calls=2,
        staged_jsonl=True,
    )
    assert len(result.understanding.items) == 3
    assert result.packet_runs[0].status == "partial"
    assert result.packet_runs[0].model_calls == 2
    stages = result.packet_runs[0].diagnostics["stages"]
    assert stages["items"]["partial_jsonl_salvaged"] is True


def test_zero_budget_is_explicitly_deferred(tmp_path: Path) -> None:
    evidence_run = build_evidence_run(_source(tmp_path), packet_chars=1000)
    llm = Mock(side_effect=AssertionError("must not call"))
    result = extract_material_understanding(evidence_run, llm=llm, max_calls=0)
    assert not result.understanding.items
    assert result.packet_runs[0].status == "deferred"
    assert result.understanding.coverage.omitted_areas
    llm.assert_not_called()


def test_slot_protocol_requires_staged_jsonl(tmp_path: Path) -> None:
    evidence_run = build_evidence_run(_source(tmp_path), packet_chars=1000)
    import pytest

    with pytest.raises(ValueError, match="requires staged_jsonl"):
        extract_material_understanding(
            evidence_run,
            llm=None,
            max_calls=0,
            slot_protocol=True,
        )


def test_service_accepts_uningested_path_without_persistence(tmp_path: Path) -> None:
    service = CorpusService("postgresql://unused")
    service._connect = Mock(side_effect=AssertionError("database must not be used"))  # type: ignore[method-assign]
    result = service.understand_material(path=_source(tmp_path), llm=_response, max_calls=1)
    assert result.understanding.source.material_type == "conference_minutes"
    assert result.summary()["items"] == 3


def test_supplied_source_is_classified_without_ingestion() -> None:
    source = Path("data/corpus/5月：锂电铜箔和电子铜箔.md")
    run = build_evidence_run(source, packet_chars=2000)
    assert classify_material_type(run.document) == "conference_minutes"


def test_unlabelled_document_voice_is_degraded_not_rejected(tmp_path: Path) -> None:
    source = tmp_path / "unlabelled.docx"
    from docx import Document

    document = Document()
    document.add_paragraph("技术演进会重构供应链格局，未来需求可能继续增长。")
    document.add_paragraph("但客户验证仍有不确定性，这并不一定代表订单已经落地。")
    document.save(source)
    evidence_run = build_evidence_run(source, packet_chars=1000)
    structure = build_material_structure(evidence_run.document)
    slots = build_candidate_slots(evidence_run.document, structure)

    assert structure.dialogue_structure == "document_voice"
    assert structure.attribution_capability == "document_only"
    assert structure.processing_mode == "degraded"
    assert slots
    assert {"forecast", "risk", "negation"} <= set(
        signal for slot in slots for signal in slot.signal_types
    )


def test_slot_protocol_recovers_null_speaker_without_cascade(tmp_path: Path) -> None:
    evidence_run = build_evidence_run(_source(tmp_path), packet_chars=1000)
    structure = build_material_structure(evidence_run.document)
    slots = build_candidate_slots(evidence_run.document, structure)
    payload = json.loads(_response(""))
    payload["speakers"] = [
        {
            "speaker_id": "summary_author",
            "display_name": None,
            "role": "summary_author",
            "identity_status": "unknown",
        }
    ]
    item = payload["items"][1]
    slot = next(slot for slot in slots if item["evidence_quote"] in slot.text)
    item["speaker_ref"] = "summary_author"
    records = [
        {"record_type": "speaker", **payload["speakers"][0]},
        {
            "record_type": "item",
            "candidate_slot_id": slot.candidate_slot_id,
            **item,
        },
        *(
            {
                "record_type": "coverage",
                "candidate_slot_id": other.candidate_slot_id,
                "status": "no_supported_item",
                "reason_code": "no_supported_item",
            }
            for other in slots
            if other.candidate_slot_id != slot.candidate_slot_id
        ),
    ]
    responses = iter(("\n".join(json.dumps(record, ensure_ascii=False) for record in records), ""))
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: next(responses),
        max_calls=2,
        staged_jsonl=True,
        slot_protocol=True,
    )

    assert len(result.understanding.items) == 1
    assert any(
        speaker.display_name is None and speaker.role == "summary_author"
        for speaker in result.understanding.speakers
    )
    assert any(entry.status == "extracted" for entry in result.understanding.coverage.slot_ledger)
    assert result.packet_runs[0].status == "partial"


def test_slot_protocol_binds_omitted_slot_id_by_unique_exact_quote(tmp_path: Path) -> None:
    evidence_run = build_evidence_run(_source(tmp_path), packet_chars=1000)
    structure = build_material_structure(evidence_run.document)
    slots = build_candidate_slots(evidence_run.document, structure)
    payload = json.loads(_response(""))
    item = payload["items"][0]
    slot = next(slot for slot in slots if item["evidence_quote"] in slot.text)
    records = [
        *({"record_type": "speaker", **speaker} for speaker in payload["speakers"]),
        {"record_type": "item", **item},
        *(
            {
                "record_type": "coverage",
                "candidate_slot_id": other.candidate_slot_id,
                "status": "no_supported_item",
                "reason_code": "checked",
            }
            for other in slots
            if other.candidate_slot_id != slot.candidate_slot_id
        ),
    ]
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: "\n".join(
            json.dumps(record, ensure_ascii=False) for record in records
        ),
        max_calls=1,
        staged_jsonl=True,
        slot_protocol=True,
    )

    assert len(result.understanding.items) == 1
    extracted_entry = next(
        entry
        for entry in result.understanding.coverage.slot_ledger
        if entry.candidate_slot_id == slot.candidate_slot_id
    )
    assert extracted_entry.status == "extracted"
    assert extracted_entry.item_refs == (result.understanding.items[0].item_id,)


def test_slot_protocol_marks_missing_terminal_record_partial(tmp_path: Path) -> None:
    evidence_run = build_evidence_run(_source(tmp_path), packet_chars=1000)
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: "",
        max_calls=2,
        staged_jsonl=True,
        slot_protocol=True,
    )

    assert result.packet_runs[0].status == "partial"
    assert result.understanding.coverage.slot_ledger[0].status == "partial"
    assert "terminal_record_missing" in result.understanding.coverage.slot_ledger[0].reason_codes
    assert result.summary()["complete"] is False


def test_exact_item_limit_is_capacity_saturation(tmp_path: Path) -> None:
    evidence_run = build_evidence_run(_source(tmp_path), packet_chars=1000)
    payload = json.loads(_response(""))
    payload["items"] = payload["items"][:1]
    item_lines = "\n".join(
        json.dumps(record, ensure_ascii=False)
        for record in [
            *({"record_type": "speaker", **speaker} for speaker in payload["speakers"]),
            {"record_type": "item", **payload["items"][0]},
        ]
    )
    responses = iter((item_lines, ""))
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: next(responses),
        max_calls=2,
        staged_jsonl=True,
        max_items_per_packet=1,
    )

    assert result.packet_runs[0].status == "partial"
    assert result.packet_runs[0].diagnostics["stages"]["items"]["item_limit_saturated"] is True


def test_slot_protocol_falls_back_to_explicit_source_metadata(tmp_path: Path) -> None:
    source = tmp_path / "James-Bulltard_9326 复盘.md"
    source.write_text("# 交易记录\n\nTTD - I bought 3000 shares today.\n", encoding="utf-8")
    evidence_run = build_evidence_run(source, packet_chars=1000)
    structure = build_material_structure(evidence_run.document)
    slots = build_candidate_slots(evidence_run.document, structure)
    slot = next(slot for slot in slots if "bought" in slot.text)
    item = {
        "record_type": "item",
        "candidate_slot_id": slot.candidate_slot_id,
        "item_id": "trade",
        "text": "Bought 3000 TTD shares",
        "semantic_type": "behavior",
        "statement_role": "claim",
        "speech_role": "statement",
        "perspective": "source_explicit",
        "speaker_ref": "undeclared",
        "polarity": "affirmed",
        "value": "3000 shares",
        "behavior_status": "claimed_executed",
        "temporal_frame": "contemporaneous",
        "evidence_quote": "TTD - I bought 3000 shares today.",
        "unknown_fields": [],
    }
    records = [
        item,
        *(
            {
                "record_type": "coverage",
                "candidate_slot_id": candidate.candidate_slot_id,
                "status": "no_supported_item",
                "reason_code": "checked",
            }
            for candidate in slots
            if candidate.candidate_slot_id != slot.candidate_slot_id
        ),
    ]
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: "\n".join(json.dumps(record, ensure_ascii=False) for record in records),
        max_calls=1,
        staged_jsonl=True,
        slot_protocol=True,
    )
    extracted = result.understanding.items[0]
    speaker = next(
        speaker
        for speaker in result.understanding.speakers
        if speaker.speaker_id == extracted.speaker_ref
    )

    assert speaker.display_name == "James-Bulltard"
    assert speaker.role == "source_author"
    assert speaker.identity_status == "explicit"
    assert extracted.perspective == "source_explicit"
    assert "speaker_reference" not in extracted.unknown_fields


def test_slot_coverage_system_corrects_detected_question_structure(tmp_path: Path) -> None:
    source = tmp_path / "qa.md"
    source.write_text("一、需求是否会改善？\n未来需求可能改善。\n", encoding="utf-8")
    evidence_run = build_evidence_run(source, packet_chars=1000)
    structure = build_material_structure(evidence_run.document)
    slot = next(
        slot
        for slot in build_candidate_slots(evidence_run.document, structure)
        if "question" in slot.signal_types
    )
    records = [
        {
            "record_type": "item",
            "candidate_slot_id": slot.candidate_slot_id,
            "item_id": "answer_only",
            "text": "错误地把问句写成预测回答",
            "semantic_type": "forecast",
            "statement_role": "answer",
            "speech_role": "answer",
            "perspective": "unknown",
            "speaker_ref": "undeclared",
            "polarity": "affirmed",
            "value": None,
            "behavior_status": None,
            "temporal_frame": "unknown",
            "evidence_quote": "需求是否会改善？",
            "unknown_fields": [],
        },
    ]
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: "\n".join(json.dumps(record, ensure_ascii=False) for record in records),
        max_calls=1,
        staged_jsonl=True,
        slot_protocol=True,
    )

    ledger = result.understanding.coverage.slot_ledger[0]
    assert ledger.status == "extracted"
    assert result.understanding.items[0].statement_role == "question"
    assert result.understanding.items[0].speech_role == "question"
    assert result.understanding.items[0].polarity == "unknown"


def test_system_owns_numbered_answers_and_quoted_continuation(tmp_path: Path) -> None:
    source = tmp_path / "qa.md"
    source.write_text(
        "八、未来行业是否值得投资？\n"
        "万物皆周期。\n"
        "专家：之前跟总工探讨过，他说工艺和设备是64开，工艺占大头。\n",
        encoding="utf-8",
    )
    evidence_run = build_evidence_run(source, packet_chars=1000)
    slots = build_candidate_slots(
        evidence_run.document, build_material_structure(evidence_run.document)
    )
    cycle = next(slot for slot in slots if "万物皆周期" in slot.text)
    process = next(slot for slot in slots if "工艺占大头" in slot.text)
    records = [
        {
            "record_type": "item",
            "candidate_slot_id": cycle.candidate_slot_id,
            "item_id": "cycle",
            "text": "万物皆周期",
            "semantic_type": "opinion",
            "statement_role": "claim",
            "speech_role": "statement",
            "perspective": "quoted_other",
            "speaker_ref": "undeclared",
            "polarity": "affirmed",
            "value": None,
            "behavior_status": None,
            "temporal_frame": "unknown",
            "evidence_quote": "万物皆周期",
            "unknown_fields": [],
        },
        {
            "record_type": "item",
            "candidate_slot_id": process.candidate_slot_id,
            "item_id": "process",
            "text": "工艺占大头",
            "semantic_type": "opinion",
            "statement_role": "claim",
            "speech_role": "statement",
            "perspective": "source_explicit",
            "speaker_ref": "undeclared",
            "polarity": "affirmed",
            "value": None,
            "behavior_status": None,
            "temporal_frame": "unknown",
            "evidence_quote": "工艺占大头",
            "unknown_fields": [],
        },
    ]
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: "\n".join(
            json.dumps(record, ensure_ascii=False) for record in records
        ),
        max_calls=1,
        material_type="research_report",
        staged_jsonl=True,
        slot_protocol=True,
        candidate_slot_ids=(cycle.candidate_slot_id, process.candidate_slot_id),
    )
    cycle_item, process_item = result.understanding.items
    speakers = {speaker.speaker_id: speaker for speaker in result.understanding.speakers}

    assert (cycle_item.statement_role, cycle_item.speech_role) == ("answer", "answer")
    assert cycle_item.perspective == "source_explicit"
    assert speakers[cycle_item.speaker_ref].role == "analyst_author"
    assert (process_item.statement_role, process_item.speech_role) == ("evidence", "answer")
    assert process_item.perspective == "quoted_other"
    assert speakers[process_item.speaker_ref].role == "quoted_source"


def test_system_owns_trade_time_and_explicit_value(tmp_path: Path) -> None:
    source = tmp_path / "James-Bulltard_9326 复盘.md"
    quote = "TTD - I added 10,000 more at 14.76."
    source.write_text(quote, encoding="utf-8")
    evidence_run = build_evidence_run(source, packet_chars=1000)
    slot = build_candidate_slots(
        evidence_run.document, build_material_structure(evidence_run.document)
    )[0]
    record = {
        "record_type": "item",
        "candidate_slot_id": slot.candidate_slot_id,
        "item_id": "trade",
        "text": "Added TTD shares",
        "semantic_type": "behavior",
        "statement_role": "claim",
        "speech_role": "statement",
        "perspective": "source_explicit",
        "speaker_ref": "undeclared",
        "polarity": "affirmed",
        "value": None,
        "behavior_status": "executed",
        "temporal_frame": "retrospective",
        "evidence_quote": quote,
        "unknown_fields": [],
    }
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: json.dumps(record, ensure_ascii=False),
        max_calls=1,
        material_type="post_trade_review",
        staged_jsonl=True,
        slot_protocol=True,
        candidate_slot_ids=(slot.candidate_slot_id,),
    )
    item = result.understanding.items[0]

    assert item.behavior_status == "claimed_executed"
    assert item.temporal_frame == "contemporaneous"
    assert item.value == "10,000 at 14.76"
    assert "external_verification" in item.unknown_fields


def test_system_owns_risk_polarity_and_report_values(tmp_path: Path) -> None:
    source = tmp_path / "report.md"
    source.write_text(
        "维持26-28年EPS预测值67.74/70.77/73.84元。\n"
        "风险提示：消费复苏不及预期。\n",
        encoding="utf-8",
    )
    evidence_run = build_evidence_run(source, packet_chars=1000)
    slots = build_candidate_slots(
        evidence_run.document, build_material_structure(evidence_run.document)
    )
    eps = next(slot for slot in slots if "EPS" in slot.text)
    risk = next(slot for slot in slots if "消费复苏" in slot.text)
    records = [
        {
            "record_type": "item",
            "candidate_slot_id": eps.candidate_slot_id,
            "item_id": "eps",
            "text": "维持EPS预测",
            "semantic_type": "forecast",
            "statement_role": "claim",
            "speech_role": "statement",
            "perspective": "source_explicit",
            "speaker_ref": "undeclared",
            "polarity": "unknown",
            "value": None,
            "behavior_status": None,
            "temporal_frame": "unknown",
            "evidence_quote": "维持26-28年EPS预测值67.74/70.77/73.84元",
            "unknown_fields": [],
        },
        {
            "record_type": "item",
            "candidate_slot_id": risk.candidate_slot_id,
            "item_id": "risk",
            "text": "消费复苏不及预期",
            "semantic_type": "forecast",
            "statement_role": "risk",
            "speech_role": "statement",
            "perspective": "source_explicit",
            "speaker_ref": "undeclared",
            "polarity": "negated",
            "value": None,
            "behavior_status": None,
            "temporal_frame": "unknown",
            "evidence_quote": "消费复苏不及预期",
            "unknown_fields": [],
        },
    ]
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: "\n".join(
            json.dumps(record, ensure_ascii=False) for record in records
        ),
        max_calls=1,
        material_type="research_report",
        staged_jsonl=True,
        slot_protocol=True,
        candidate_slot_ids=(eps.candidate_slot_id, risk.candidate_slot_id),
    )
    eps_item, risk_item = result.understanding.items

    assert eps_item.value == "26-28年 67.74/70.77/73.84元"
    assert eps_item.polarity == "affirmed"
    assert risk_item.polarity == "affirmed"


def test_system_preserves_claim_polarity_when_text_contains_lexical_negation(
    tmp_path: Path,
) -> None:
    source = tmp_path / "report.md"
    quote = "茅台经营向上明确，底层逻辑未变"
    source.write_text(quote, encoding="utf-8")
    evidence_run = build_evidence_run(source, packet_chars=1000)
    slot = build_candidate_slots(
        evidence_run.document, build_material_structure(evidence_run.document)
    )[0]
    record = {
        "record_type": "item",
        "candidate_slot_id": slot.candidate_slot_id,
        "item_id": "outlook",
        "text": quote,
        "semantic_type": "opinion",
        "statement_role": "claim",
        "speech_role": "statement",
        "perspective": "source_explicit",
        "speaker_ref": "undeclared",
        "polarity": "affirmed",
        "value": None,
        "behavior_status": None,
        "temporal_frame": "contemporaneous",
        "evidence_quote": quote,
        "unknown_fields": [],
    }
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: json.dumps(record, ensure_ascii=False),
        max_calls=1,
        material_type="research_report",
        staged_jsonl=True,
        slot_protocol=True,
        candidate_slot_ids=(slot.candidate_slot_id,),
    )

    assert result.understanding.items[0].polarity == "affirmed"


def test_system_adds_controlled_unknown_axes_for_unattributed_summary(tmp_path: Path) -> None:
    source = tmp_path / "minutes.md"
    source.write_text("总体而言，行业仍有增长潜力。", encoding="utf-8")
    evidence_run = build_evidence_run(source, packet_chars=1000)
    slot = build_candidate_slots(
        evidence_run.document, build_material_structure(evidence_run.document)
    )[0]
    record = {
        "record_type": "item",
        "candidate_slot_id": slot.candidate_slot_id,
        "item_id": "summary",
        "text": "行业仍有增长潜力",
        "semantic_type": "opinion",
        "statement_role": "claim",
        "speech_role": "statement",
        "perspective": "source_explicit",
        "speaker_ref": "undeclared",
        "polarity": "affirmed",
        "value": None,
        "behavior_status": None,
        "temporal_frame": "unknown",
        "evidence_quote": "总体而言，行业仍有增长潜力",
        "unknown_fields": [],
    }
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: json.dumps(record, ensure_ascii=False),
        max_calls=1,
        material_type="conference_minutes",
        staged_jsonl=True,
        slot_protocol=True,
        candidate_slot_ids=(slot.candidate_slot_id,),
    )
    item = result.understanding.items[0]

    assert item.perspective == "unknown"
    assert {
        "identity",
        "time",
        "value",
        "summary_authorship",
        "summary_generation_method",
    } <= set(item.unknown_fields)


def test_system_splits_known_compound_statements_into_atomic_obligations(
    tmp_path: Path,
) -> None:
    source = tmp_path / "atomic.md"
    source.write_text(
        "维持一年目标价2030元和“强推”评级。\n"
        "风险提示：宏观需求持续疲软、消费复苏不及预期、行业竞争加剧。\n"
        "TTD - I added 10,000 more at 14.76 and sold 100 $15 calls for .16.\n"
        "专家：设备精度已经做到行业领先，那么大概率跟工艺挂钩。\n"
        "专家：他说工艺和设备是64开，工艺占大头，设备只是配套，但没设备也不行。\n",
        encoding="utf-8",
    )
    evidence_run = build_evidence_run(source, packet_chars=2000)
    structure = build_material_structure(evidence_run.document)
    slots = build_candidate_slots(evidence_run.document, structure)
    slot_texts = ["".join(slot.text.split()) for slot in slots]

    expected_atoms = (
        "维持一年目标价2030元",
        "和“强推”评级",
        "宏观需求持续疲软",
        "消费复苏不及预期",
        "行业竞争加剧",
        "TTD-Iadded10,000moreat14.76",
        "andsold100$15callsfor.16",
        "设备精度已经做到行业领先",
        "那么大概率跟工艺挂钩",
        "他说工艺和设备是64开",
        "工艺占大头",
        "设备只是配套",
        "但没设备也不行",
    )
    for atom in expected_atoms:
        assert sum(atom in text for text in slot_texts) == 1
    assert len(slots) == len(expected_atoms)


def test_system_preserves_connector_commas_inside_one_question(tmp_path: Path) -> None:
    source = tmp_path / "question.md"
    question = "未来行业两年供需很好，但更长期可能扩产，是否值得投资？"
    source.write_text(question, encoding="utf-8")
    evidence_run = build_evidence_run(source, packet_chars=1000)
    slots = build_candidate_slots(
        evidence_run.document, build_material_structure(evidence_run.document)
    )

    assert len(slots) == 1
    assert slots[0].text == question
    assert slots[0].signal_types[0] == "question"


def test_system_does_not_create_obligations_for_bare_acknowledgements(tmp_path: Path) -> None:
    source = tmp_path / "acknowledgements.md"
    source.write_text("投资者：明白。\n专家：好的！\n", encoding="utf-8")
    evidence_run = build_evidence_run(source, packet_chars=1000)
    slots = build_candidate_slots(
        evidence_run.document, build_material_structure(evidence_run.document)
    )

    assert slots == ()


def test_lexical_signals_allow_supported_semantic_disambiguation(tmp_path: Path) -> None:
    source = tmp_path / "signals.md"
    source.write_text(
        "关键仍是对未来景气趋势的判断。\n"
        "the CEO bought $150m worth\n",
        encoding="utf-8",
    )
    evidence_run = build_evidence_run(source, packet_chars=1000)
    slots = build_candidate_slots(
        evidence_run.document, build_material_structure(evidence_run.document)
    )
    forecast_slot = next(slot for slot in slots if "未来景气" in slot.text)
    behavior_slot = next(slot for slot in slots if "CEO bought" in slot.text)
    records = [
        {
            "record_type": "item",
            "candidate_slot_id": forecast_slot.candidate_slot_id,
            "item_id": "future_opinion",
            "text": "未来景气趋势仍需判断",
            "semantic_type": "opinion",
            "statement_role": "claim",
            "speech_role": "statement",
            "perspective": "source_explicit",
            "speaker_ref": "undeclared",
            "polarity": "affirmed",
            "value": None,
            "behavior_status": None,
            "temporal_frame": "unknown",
            "evidence_quote": "关键仍是对未来景气趋势的判断。",
            "unknown_fields": [],
        },
        {
            "record_type": "item",
            "candidate_slot_id": behavior_slot.candidate_slot_id,
            "item_id": "third_party_trade",
            "text": "CEO bought $150m worth",
            "semantic_type": "fact",
            "statement_role": "claim",
            "speech_role": "statement",
            "perspective": "source_explicit",
            "speaker_ref": "undeclared",
            "polarity": "affirmed",
            "value": "$150m",
            "behavior_status": None,
            "temporal_frame": "unknown",
            "evidence_quote": "the CEO bought $150m worth",
            "unknown_fields": [],
        },
    ]
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: "\n".join(
            json.dumps(record, ensure_ascii=False) for record in records
        ),
        max_calls=1,
        staged_jsonl=True,
        slot_protocol=True,
        max_slots_per_batch=2,
    )

    assert [entry.status for entry in result.understanding.coverage.slot_ledger] == [
        "extracted",
        "extracted",
    ]


def test_mixed_audio_turn_is_one_unknown_attribution_obligation(tmp_path: Path) -> None:
    source = tmp_path / "mixed.md"
    source.write_text("投资者：大家好，请问能听到吗？可以，您讲。\n", encoding="utf-8")
    evidence_run = build_evidence_run(source, packet_chars=1000)
    structure = build_material_structure(evidence_run.document)
    slots = build_candidate_slots(evidence_run.document, structure)

    assert len(slots) == 1
    assert slots[0].attribution_capability == "unavailable"
    assert slots[0].explicit_role is None
    assert "question" in slots[0].signal_types


def test_atomic_obligation_batches_never_exceed_item_capacity(tmp_path: Path) -> None:
    source = tmp_path / "many-turns.md"
    source.write_text(
        "\n".join(f"投资者：第{index}个问题怎么看？" for index in range(35)),
        encoding="utf-8",
    )
    evidence_run = build_evidence_run(source, packet_chars=10_000)
    structure = build_material_structure(evidence_run.document)
    slots = build_candidate_slots(evidence_run.document, structure)
    batches = build_candidate_slot_batches(
        slots,
        max_slots_per_batch=8,
        max_items_per_batch=8,
    )

    assert len(slots) == 35
    assert len(batches) == 5
    assert all(1 <= len(batch) <= 8 for batch in batches)
    assert all(len({slot.packet_id for slot in batch}) == 1 for batch in batches)


def test_slot_protocol_calls_model_once_per_finite_atomic_batch(tmp_path: Path) -> None:
    source = tmp_path / "many-claims.md"
    source.write_text(
        "\n".join(f"这是第{index}个普通材料陈述片段。" for index in range(17)),
        encoding="utf-8",
    )
    prompts: list[str] = []

    def llm(prompt: str) -> str:
        prompts.append(prompt)
        slot_payload = prompt.split("候选槽位：\n", 1)[1].split("\n来源语境", 1)[0]
        slots = json.loads(slot_payload)
        return "\n".join(
            json.dumps(
                {
                    "record_type": "coverage",
                    "candidate_slot_id": slot["candidate_slot_id"],
                    "status": "no_supported_item",
                    "reason_code": "not_research_relevant",
                },
                ensure_ascii=False,
            )
            for slot in slots
        )

    evidence_run = build_evidence_run(source, packet_chars=10_000)
    result = extract_material_understanding(
        evidence_run,
        llm=llm,
        max_calls=3,
        staged_jsonl=True,
        slot_protocol=True,
        max_slots_per_batch=8,
        max_items_per_packet=30,
    )

    assert len(prompts) == 3
    assert result.packet_runs[0].model_calls == 3
    assert result.packet_runs[0].status == "completed"
    assert len(result.understanding.coverage.slot_ledger) == 17
    assert result.summary()["complete"] is True


def test_detected_question_cannot_be_rejected_as_no_supported_item(tmp_path: Path) -> None:
    source = tmp_path / "question.md"
    source.write_text("未来需求会改善吗？\n", encoding="utf-8")
    evidence_run = build_evidence_run(source, packet_chars=1000)
    slot = build_candidate_slots(
        evidence_run.document, build_material_structure(evidence_run.document)
    )[0]
    response = json.dumps(
        {
            "record_type": "coverage",
            "candidate_slot_id": slot.candidate_slot_id,
            "status": "no_supported_item",
            "reason_code": "model_rejected",
        },
        ensure_ascii=False,
    )
    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: response,
        max_calls=1,
        staged_jsonl=True,
        slot_protocol=True,
    )

    ledger = result.understanding.coverage.slot_ledger[0]
    assert ledger.status == "partial"
    assert "detected_signal_rejected:question" in ledger.reason_codes
    assert result.summary()["complete"] is False


def test_relation_candidates_exclude_adjacent_quote_without_explicit_connector() -> None:
    def item(
        item_id: str,
        quote: str,
        *,
        semantic_type: str,
        statement_role: str,
        perspective: str,
    ) -> MaterialItem:
        return MaterialItem(
            item_id=item_id,
            text=quote,
            semantic_type=semantic_type,  # type: ignore[arg-type]
            statement_role=statement_role,  # type: ignore[arg-type]
            speech_role="statement",
            perspective=perspective,  # type: ignore[arg-type]
            speaker_ref="speaker",
            polarity="affirmed",
            temporal_frame="contemporaneous",
            evidence=(
                MaterialEvidence(
                    source_rev="source",
                    packet_id="packet",
                    locator="document",
                    quote=quote,
                    start=0,
                    end=len(quote),
                ),
            ),
        )

    behavior = item(
        "behavior",
        "I added 10,000 more at 14.76",
        semantic_type="behavior",
        statement_role="claim",
        perspective="source_explicit",
    )
    adjacent_ceo_fact = item(
        "ceo",
        "the CEO bought $150m worth",
        semantic_type="fact",
        statement_role="evidence",
        perspective="quoted_other",
    )
    explicit_quote = item(
        "quote",
        "他说工艺和设备是64开",
        semantic_type="opinion",
        statement_role="evidence",
        perspective="quoted_other",
    )

    assert _relation_candidate_pairs([behavior, adjacent_ceo_fact]) == []
    pairs = _relation_candidate_pairs([behavior, explicit_quote])
    assert len(pairs) == 1
    assert pairs[0] == {
        "from_item": "quote",
        "to_item": "behavior",
        "allowed_type": "supports",
        "candidate_pair_id": pairs[0]["candidate_pair_id"],
    }


def test_relation_obligations_require_one_explicit_decision_per_pair(tmp_path: Path) -> None:
    source = tmp_path / "relation.md"
    source.write_text("结论成立。因为数据改善。\n", encoding="utf-8")
    packet = build_evidence_run(source, packet_chars=1000).document.packets[0]
    pair = {
        "candidate_pair_id": "pair_one",
        "from_item": "evidence",
        "to_item": "claim",
        "allowed_type": "supports",
    }

    relations, incomplete, counts = _relations_from_decisions(
        [
            {
                "record_type": "relation_decision",
                "candidate_pair_id": "pair_one",
                "status": "present",
                "evidence_quote": "因为数据改善",
            }
        ],
        packet,
        "source",
        [pair],
    )
    assert incomplete is False
    assert len(relations) == 1
    assert counts["decisions"] == 1

    relations, incomplete, counts = _relations_from_decisions([], packet, "source", [pair])
    assert relations == []
    assert incomplete is True
    assert counts["missing_decisions"] == 1


def test_candidate_slot_scope_runs_only_frozen_obligations(tmp_path: Path) -> None:
    source = tmp_path / "scoped.md"
    source.write_text("第一项普通判断。第二项普通判断。第三项普通判断。\n", encoding="utf-8")
    evidence_run = build_evidence_run(source, packet_chars=1000)
    slots = build_candidate_slots(
        evidence_run.document, build_material_structure(evidence_run.document)
    )
    selected = slots[1]

    result = extract_material_understanding(
        evidence_run,
        llm=lambda _prompt: json.dumps(
            {
                "record_type": "coverage",
                "candidate_slot_id": selected.candidate_slot_id,
                "status": "no_supported_item",
                "reason_code": "not_research_relevant",
            },
            ensure_ascii=False,
        ),
        max_calls=1,
        staged_jsonl=True,
        slot_protocol=True,
        extract_relations=False,
        candidate_slot_ids=(selected.candidate_slot_id,),
    )

    assert result.candidate_slots == (selected,)
    assert len(result.understanding.coverage.slot_ledger) == 1
    assert (
        result.understanding.coverage.slot_ledger[0].candidate_slot_id == selected.candidate_slot_id
    )
