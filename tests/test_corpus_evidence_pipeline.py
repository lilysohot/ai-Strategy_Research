"""End-to-end evidence contracts using synthetic sources, not copied private reports."""

from dataclasses import replace
from decimal import Decimal

import pymupdf
import pytest
from docx import Document

from plugins.corpus.derivation import derive
from plugins.corpus.evidence import fingerprint, parse_evidence, split_spans
from plugins.corpus.evidence_pipeline import EvidenceRun, build_evidence_run
from plugins.corpus.ingest import parse_docx, parse_markdown


def make_report(tmp_path):
    path = tmp_path / "2026-08-16_600519.SH.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    # ASCII source still exercises real PDF coordinates and multilingual independent parsing.
    for x, text in [(40, "Metric"), (200, "2025A"), (280, "2026E")]:
        page.insert_text((x, 100), text)
    for y, label, a, b in [(120, "revenue", "100", "120"), (140, "Zero", "0", "-25")]:
        for x, text in [(40, label), (200, a), (280, b)]:
            page.insert_text((x, y), text)
    page.insert_text((40, 50), "600519.SH financial forecast")
    pdf.save(path)
    pdf.close()
    return path


def test_pdf_cells_have_verifiable_coordinates_and_raw_zero(tmp_path):
    doc = parse_evidence(make_report(tmp_path))
    cells = [c for p in doc.packets for c in p.cells]
    assert {(c.row, c.column, c.value) for c in cells} == {
        ("revenue", "2025A", "100"),
        ("revenue", "2026E", "120"),
        ("Zero", "2025A", "0"),
        ("Zero", "2026E", "-25"),
    }
    assert all(c.span.bbox is not None and c.column_span.bbox is not None for c in cells)
    packet = next(p for p in doc.packets if p.cells)
    good = packet.cells[0].reference()
    assert packet.lookup(good) == "100"
    assert packet.lookup({**good, "column": "2027E"}) is None
    assert packet.lookup({**good, "cell": "10"}) is None
    assert "-25" in doc.pages[0].text


def test_long_text_coverage_and_parser_configuration_identity(tmp_path):
    text = "A line with 100 and a condition.\n" * 300
    spans = split_spans(text, "1", 200)
    assert "".join(s.text for s in spans) == text
    assert max(len(s.text) for s in spans) <= 200
    assert spans[0].start == 0 and spans[-1].end == len(text)
    path = make_report(tmp_path)
    first = parse_evidence(path, packet_chars=200)
    same = parse_evidence(path, packet_chars=200)
    changed = parse_evidence(path, packet_chars=300)
    assert first.parse_rev == same.parse_rev
    assert first.source_rev == changed.source_rev
    assert first.parse_rev != changed.parse_rev


def test_image_only_page_is_unknown_not_empty_success(tmp_path):
    path = tmp_path / "scan.pdf"
    pdf = pymupdf.open()
    pdf.new_page()
    pdf.save(path)
    pdf.close()
    run = build_evidence_run(path)
    assert run.summary()["complete"] is False
    assert run.packet_runs[0].status == "unknown"
    assert "needs_ocr" in run.packet_runs[0].reasons


def test_markdown_headings_and_word_tables_are_preserved(tmp_path):
    path = tmp_path / "source.md"
    path.write_text("# Forecast\n## Revenue\n100\n", encoding="utf-8")
    assert "## Revenue" in parse_markdown(path)[-1].text
    doc = Document()
    doc.add_heading("Forecast", level=1)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Revenue", "2026E"
    table.cell(1, 0).text, table.cell(1, 1).text = "Amount", "100"
    target = tmp_path / "source.docx"
    doc.save(target)
    text = "\n".join(b.text for b in parse_docx(target))
    assert "Forecast" in text and "Amount | 100" in text


def test_run_roundtrip_and_no_unit_means_not_calculable(tmp_path):
    path = make_report(tmp_path)
    run = build_evidence_run(path)
    same = build_evidence_run(path)
    assert run.run_id == same.run_id
    restored = EvidenceRun.model_validate_json(run.model_dump_json())
    restored.verify_identity()
    assert all("calculate" not in f.usable_for for f in restored.facts)
    changed = restored.model_copy(update={"model": "tampered"})
    with pytest.raises(ValueError, match="hash mismatch"):
        changed.verify_identity()


def _calculation_run(tmp_path):
    # Change only units via a real evidence source, not by trusting model-provided table_ref.
    path = tmp_path / "2026-08-16_600519.SH.md"
    path.write_text("# 600519.SH\n2025A 营业收入100元。2026E 营业收入120元。", encoding="utf-8")
    import json

    payload = [
        dict(
            claim_text=f"{period} 营业收入{value}元",
            evidence_quote=f"{period} 营业收入{value}元",
            scope="company",
            subject="600519.SH",
            metric="营业收入",
            value_text=f"{value}元",
            period_raw=period,
            kind=kind,
        )
        for period, value, kind in [("2025A", "100", "fact"), ("2026E", "120", "forecast")]
    ]
    return build_evidence_run(
        path, llm=lambda _: json.dumps(payload, ensure_ascii=False), model="test", max_prose_calls=3
    )


def test_growth_has_exact_input_ids_and_refuses_reversed_periods(tmp_path):
    run = _calculation_run(tmp_path)
    assert len(run.facts) == 2
    by_period = {f.claim.period_raw: f for f in run.facts}
    inputs = (by_period["2026E"].fact_id, by_period["2025A"].fact_id)
    result = derive(run, "revenue_growth", inputs)
    assert result.value == Decimal("20")
    assert result.inputs == inputs and result.basis == "source_forecast"
    with pytest.raises(ValueError, match="consecutive"):
        derive(run, "revenue_growth", inputs[::-1])
    with pytest.raises(ValueError, match="duplicate"):
        derive(run, "revenue_growth", (inputs[0], inputs[0]))


def test_llm_failure_does_not_report_complete(tmp_path):
    path = tmp_path / "source.md"
    path.write_text("Revenue 2026 is 100.", encoding="utf-8")
    run = build_evidence_run(path, llm=lambda _: "[{broken", max_prose_calls=1)
    assert not run.summary()["complete"]
    assert not run.facts


def test_changed_subject_is_refused_even_with_valid_amount(tmp_path):
    run = _calculation_run(tmp_path)
    first, second = run.facts
    modified = second.model_copy(update={"claim": replace(second.claim, subject="000001.SZ")})
    invalid = run.model_copy(update={"facts": (first, modified)})
    payload = invalid.model_dump(mode="json")
    payload.pop("run_id")
    invalid = invalid.model_copy(update={"run_id": fingerprint(payload)})
    with pytest.raises(ValueError, match="subject_mismatch"):
        derive(invalid, "revenue_growth", (modified.fact_id, first.fact_id))


@pytest.mark.parametrize("response", ["", "   ", "[{broken"])
def test_empty_or_broken_response_is_failure_not_empty_success(tmp_path, response):
    path = tmp_path / "source.md"
    path.write_text("Revenue 2026 is 100.", encoding="utf-8")
    run = build_evidence_run(path, llm=lambda _: response, max_prose_calls=1)
    assert run.packet_runs[0].status == "failed"
    assert not run.summary()["complete"]


def test_source_view_does_not_apply_legacy_boilerplate_cleaning(tmp_path):
    path = tmp_path / "source.md"
    content = "# Parent\r\n## Child\r\nRevenue 100\r\n1 / 2\r\n"
    path.write_bytes(content.encode())
    doc = parse_evidence(path)
    assert doc.pages[0].text == content
    assert "".join(p.text for p in doc.packets) == content
    assert "Parent" in doc.packets[0].context and "Child" in doc.packets[0].context


def test_database_versions_idempotence_and_backup_roundtrip(tmp_path):
    import uuid
    from urllib.parse import quote

    import psycopg

    from plugins.corpus.evidence_pipeline import extract_evidence
    from plugins.corpus.service import CorpusService, dsn

    admin = dsn()
    try:
        psycopg.connect(admin, connect_timeout=5).close()
    except psycopg.OperationalError:
        pytest.skip("PostgreSQL unavailable")
    schemas = [f"corpus_evidence_test_{uuid.uuid4().hex[:12]}" for _ in range(2)]
    urls = [
        admin
        + ("&" if "?" in admin else "?")
        + "options="
        + quote(f"-c search_path={schema},public")
        for schema in schemas
    ]
    with psycopg.connect(admin, autocommit=True) as conn:
        for schema in schemas:
            conn.execute(psycopg.sql.SQL("CREATE SCHEMA {}").format(psycopg.sql.Identifier(schema)))
    try:
        service = CorpusService(urls[0])
        service.init_db()
        path = make_report(tmp_path)
        first = service.extract_claims(path)
        second = extract_evidence(parse_evidence(path, packet_chars=300))
        for run in (first, first, second):
            service.save_evidence_run(run)
            assert service.load_evidence_run(run.run_id) == run
            projection = service.claims_of(
                run_id=run.run_id,
                purpose="audit",
                quality_status=None,
            )
            assert projection["total"] == len(run.facts)
            assert projection["run_id"] == run.run_id
        assert len(service.claim_runs(doc_id=first.document.doc_id)) == 2
        packet = first.document.packets[0]
        assert service.fetch_evidence(first.run_id, packet.packet_id) == packet.model_dump(
            mode="json"
        )
        with psycopg.connect(urls[0]) as conn:
            assert conn.execute("SELECT count(*) FROM corpus_evidence_runs").fetchone()[0] == 2
            assert conn.execute("SELECT count(*) FROM documents").fetchone()[0] == 0
        backup = service.backup(tmp_path / "backup", mode="csv")
        counts = service.restore(backup, urls[1])
        assert counts["corpus_evidence_runs"] == 2
        assert CorpusService(urls[1]).load_evidence_run(first.run_id) == first
    finally:
        with psycopg.connect(admin, autocommit=True) as conn:
            for schema in schemas:
                conn.execute(
                    psycopg.sql.SQL("DROP SCHEMA {} CASCADE").format(psycopg.sql.Identifier(schema))
                )


@pytest.mark.parametrize(
    "change,reason",
    [
        ({"unit": "美元"}, "unit_mismatch"),
        ({"qualifiers": {"basis": "MoM"}}, "basis_mismatch"),
        ({"source_rev": "another-source"}, "source_revision_mismatch"),
    ],
)
def test_calculation_rejects_incompatible_inputs(tmp_path, change, reason):
    run = _calculation_run(tmp_path)
    first, second = run.facts
    second = second.model_copy(update={"claim": replace(second.claim, **change)})
    invalid = run.model_copy(update={"facts": (first, second)})
    payload = invalid.model_dump(mode="json")
    payload.pop("run_id")
    invalid = invalid.model_copy(update={"run_id": fingerprint(payload)})
    with pytest.raises(ValueError, match=reason):
        derive(invalid, "revenue_growth", (second.fact_id, first.fact_id))


def test_wrong_unit_or_metric_cannot_borrow_a_real_number(tmp_path):
    import json

    path = tmp_path / "2026-08-16_600519.SH.md"
    path.write_text("2026E 营业收入100亿元。", encoding="utf-8")
    for metric, unit, reason in [
        ("营业收入", "元", "value_unit_pair_not_in_quote"),
        ("净利润", "亿元", "metric_not_in_quote"),
    ]:
        payload = [
            {
                "claim_text": f"2026E {metric}100{unit}",
                "evidence_quote": "2026E 营业收入100亿元",
                "subject": "600519.SH",
                "metric": metric,
                "value_text": "100",
                "unit_raw": unit,
                "period_raw": "2026E",
                "kind": "forecast",
            }
        ]
        run = build_evidence_run(
            path, llm=lambda _, payload=payload: json.dumps(payload), max_prose_calls=1
        )
        assert reason in run.facts[0].reasons
        assert "calculate" not in run.facts[0].usable_for
