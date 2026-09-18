"""Bounded follow-up probes for the previously reported F1/F2/F4/F7 contracts.

Synthetic inputs only. Original acceptance probes remain unchanged.
"""
import io
import json
from dataclasses import replace
from datetime import UTC, datetime, timedelta

import pymupdf
import pytest
from docx import Document

from plugins.corpus.preparation.admission import AdmissionError, load_admission_policy
from plugins.corpus.preparation.chunk import chunk_clean_result
from plugins.corpus.preparation.clean import clean_reader_result
from plugins.corpus.preparation.contract import (
    Admission, AdmissionDecision, Build, JobStage, LeaseConfig, MaterialType,
    ResearchDomain, Unit, compute_build_id, source_id_from_bytes,
)
from plugins.corpus.preparation.readers import read_document
from plugins.corpus.preparation.repository import MemoryStore, StoreError

NOW = datetime(2026, 9, 16, tzinfo=UTC)
SID = "a" * 64
LEASE = LeaseConfig(120, 30, 360, 3)


def setup_store():
    store = MemoryStore()
    admission = Admission("d0", SID, MaterialType.RESEARCH_REPORT,
                          ResearchDomain.COMPANY, AdmissionDecision.IN_SCOPE)
    args = dict(source_id=SID, decision_id="d0", parse_rev="p1", clean_rev="c1",
                chunk_rev="k1", index_rev="i1", scope_ref=None)
    build = Build(build_id=compute_build_id(**args), **args)
    store.put_admission(admission)
    store.put_build(build)
    return store, admission, build


def unit_for(build):
    return Unit("u0", build.build_id, "paragraph", "value 12", source_id_from_bytes(b"value 12"))


def test_f1_publish_checks_build_decision_even_when_scope_unchanged():
    store, admission, build = setup_store()
    store.put_admission(replace(admission, decision_id="d1"))
    with pytest.raises(StoreError):
        store.publish(SID, "d1", build.build_id, NOW)


def test_f1_same_publish_retry_is_idempotent():
    store, _, build = setup_store()
    # C1 fail-closed ownership gate: publish goes through the engine job protocol
    # (protocol upgrade only; the idempotency assertion is unchanged). Acquire uses
    # the real wall clock because MemoryStore judges lease expiry against real now.
    store.register_job(build.build_id, JobStage.PUBLISHED)
    acquired = store.acquire_job(build.build_id, JobStage.PUBLISHED, "fixture",
                                 datetime.now(UTC), LEASE)
    token = acquired.fence_token
    assert token is not None
    first = store.publish(SID, "d0", build.build_id, NOW,
                          owner_id="fixture", fence_token=token)
    second = store.publish(SID, "d0", build.build_id, NOW,
                           owner_id="fixture", fence_token=token)
    assert second == first


def test_f2_expired_token_cannot_write_without_takeover():
    store, _, build = setup_store()
    store.register_job(build.build_id, JobStage.PARSED)
    job = store.acquire_job(build.build_id, JobStage.PARSED, "a", NOW, LEASE)
    # The same job is known expired through the public heartbeat interface.
    with pytest.raises(StoreError, match="过期"):
        store.heartbeat_job(build.build_id, JobStage.PARSED, "a", job.fence_token,
                            NOW + timedelta(seconds=121), LEASE)
    with pytest.raises(StoreError):
        store.put_units(build.build_id, [unit_for(build)], owner_id="a", fence_token=job.fence_token)


def test_f2_missing_job_is_fail_closed_for_authoritative_writes():
    store, _, build = setup_store()
    with pytest.raises(StoreError):
        store.put_units(build.build_id, [unit_for(build)])


def test_f2_publish_has_same_ownership_gate_as_other_writes():
    store, _, build = setup_store()
    store.register_job(build.build_id, JobStage.PUBLISHED)
    store.acquire_job(build.build_id, JobStage.PUBLISHED, "a", NOW, LEASE)
    store.acquire_job(build.build_id, JobStage.PUBLISHED, "b", NOW + timedelta(seconds=121), LEASE)
    with pytest.raises(StoreError):
        store.publish(SID, "d0", build.build_id, NOW + timedelta(seconds=122))


def test_f4_medium_image_region_does_not_disappear(tmp_path):
    path = tmp_path / "medium-image.pdf"
    with pymupdf.open() as doc:
        page = doc.new_page(width=600, height=800)
        page.insert_text((72, 72), "Report with chart", fontsize=11)
        page.insert_image(pymupdf.Rect(72, 120, 472, 320),
                          pixmap=pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 10, 10)))
        doc.save(str(path))
    result = read_document(path)
    assert result.issues or any(u.kind == "image" or u.status.value != "kept" for u in result.units)


def test_f4_docx_cell_image_has_gap_record(tmp_path):
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Chart caption"
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 10, 10))
    cell.paragraphs[0].add_run().add_picture(io.BytesIO(pixmap.tobytes("png")))
    path = tmp_path / "cell-image.docx"
    doc.save(str(path))
    result = read_document(path)
    assert result.issues or any(u.kind == "image" or u.status.value != "kept" for u in result.units)


def test_f7_omitted_large_header_still_has_reference_from_continuation(tmp_path):
    path = tmp_path / "long-header.md"
    path.write_text("| " + "H" * 990 + " |\n|---|\n| " + "A" * 990 + " |\n", encoding="utf-8")
    reader = read_document(path)
    result = chunk_clean_result(reader, clean_reader_result(reader))
    header = reader.units[0]
    continuation = next(c for c in result.chunks if "A" * 990 in c.search_text)
    # An indirect context reference is sufficient; copying the full header is not required.
    assert (header.ordinal in continuation.unit_ordinals
            or continuation.title_text == header.raw_text
            or getattr(continuation, "context_refs", ())), continuation


def test_policy_rejects_unsupported_rule_revision(tmp_path):
    path = tmp_path / "unsupported-policy.json"
    path.write_text(json.dumps({"policy_rev": "v999-not-implemented", "frozen": True,
                                "auto_decision": {"enabled": False}}), encoding="utf-8")
    with pytest.raises(AdmissionError):
        load_admission_policy(path)


def test_control_wrong_token_is_refused():
    store, _, build = setup_store()
    store.register_job(build.build_id, JobStage.PARSED)
    store.acquire_job(build.build_id, JobStage.PARSED, "a", NOW, LEASE)
    with pytest.raises(StoreError):
        store.put_units(build.build_id, [unit_for(build)], owner_id="a", fence_token="wrong")


def test_control_current_owner_can_write():
    store, _, build = setup_store()
    store.register_job(build.build_id, JobStage.PARSED)
    job = store.acquire_job(build.build_id, JobStage.PARSED, "a", datetime.now(UTC), LEASE)
    unit = unit_for(build)
    store.put_units(build.build_id, [unit], owner_id="a", fence_token=job.fence_token)
    assert store.get_unit(build.build_id, unit.unit_id) == unit
