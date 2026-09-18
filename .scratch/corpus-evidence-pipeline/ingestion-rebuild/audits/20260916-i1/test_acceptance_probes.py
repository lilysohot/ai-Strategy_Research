"""Read-only I1 acceptance probes. Synthetic inputs only; intentionally red before fixes.

Run with the I1 guard before collection, --noconftest, and plugin autoload disabled.
No production modules are patched and no corpus source is opened.
"""

from dataclasses import replace
from datetime import UTC, datetime, timedelta

import pytest
import pymupdf
from docx import Document

from plugins.corpus.preparation.admission import (
    POLICY_REV_V1, AdmissionPolicy, ProbeInput, SourceDescriptor,
    decide_admission, probe_features,
)
from plugins.corpus.preparation.chunk import chunk_clean_result
from plugins.corpus.preparation.clean import CleanError, clean_reader_result, verify_clean_region
from plugins.corpus.preparation.contract import (
    Admission, AdmissionDecision, Build, DocumentFormat, JobStage, LeaseConfig,
    MaterialType, ResearchDomain, ReviewDecision, ReviewedDecision, Unit,
    compute_build_id, source_id_from_bytes,
)
from plugins.corpus.preparation.readers import read_document
from plugins.corpus.preparation.repository import MemoryStore, StoreError

NOW = datetime(2026, 9, 16, tzinfo=UTC)
SID = "a" * 64


def setup_store():
    store = MemoryStore()
    admission = Admission("d0", SID, MaterialType.RESEARCH_REPORT,
                          ResearchDomain.COMPANY, AdmissionDecision.IN_SCOPE)
    args = dict(source_id=SID, decision_id="d0", parse_rev="p1", clean_rev="c1",
                chunk_rev="k1", index_rev="i1", scope_ref=None)
    build = Build(build_id=compute_build_id(**args), **args)
    store.put_admission(admission)
    store.put_build(build)
    return store, admission, build


def test_old_admission_replay_must_not_rollback_current():
    store, old, _ = setup_store()
    revoked = replace(old, decision_id="d1", decision=AdmissionDecision.EXCLUDED_BY_POLICY)
    store.put_admission(revoked)
    store.put_admission(old)
    assert store.latest_admission(SID).decision_id == "d1"


def test_revoked_source_cannot_be_republished_by_old_decision():
    store, old, build = setup_store()
    store.put_admission(replace(old, decision_id="d1", decision=AdmissionDecision.EXCLUDED_BY_POLICY))
    store.retire(SID, "d1", NOW)
    with pytest.raises(StoreError):
        store.publish(SID, "d0", build.build_id, NOW)


def test_scoped_approval_cannot_publish_full_scope_build():
    store, old, build = setup_store()
    store.put_admission(replace(old, decision_id="d1", scope_ref="page:1"))
    with pytest.raises(StoreError):
        store.publish(SID, "d1", build.build_id, NOW)


def test_expired_running_jobs_respect_max_attempts():
    store, _, build = setup_store()
    lease = LeaseConfig(120, 30, 360, 2)
    store.register_job(build.build_id, JobStage.PARSED)
    store.acquire_job(build.build_id, JobStage.PARSED, "a", NOW, lease)
    store.acquire_job(build.build_id, JobStage.PARSED, "b", NOW + timedelta(seconds=121), lease)
    with pytest.raises(StoreError):
        store.acquire_job(build.build_id, JobStage.PARSED, "c", NOW + timedelta(seconds=242), lease)


def test_authoritative_output_cannot_be_written_without_current_ownership():
    store, _, build = setup_store()
    lease = LeaseConfig(120, 30, 360, 2)
    store.register_job(build.build_id, JobStage.PARSED)
    store.acquire_job(build.build_id, JobStage.PARSED, "a", NOW, lease)
    store.acquire_job(build.build_id, JobStage.PARSED, "b", NOW + timedelta(seconds=121), lease)
    unit = Unit("late-a", build.build_id, "paragraph", "late output",
                source_id_from_bytes(b"late output"))
    # The public output-write seam has no owner/token arguments; an old worker can call it.
    with pytest.raises(StoreError):
        store.put_units(build.build_id, [unit])


def decide_with(reviews):
    source = SourceDescriptor(SID, DocumentFormat.MARKDOWN, "synthetic.md",
                              "证券公司研究", ResearchDomain.COMPANY)
    probes = probe_features(ProbeInput(source.title, "证券研究所", ("正文。",), ()))
    policy = AdmissionPolicy(POLICY_REV_V1, tuple(ResearchDomain))
    return decide_admission(source, tuple(reviews), probes, policy)


def review(name, decision=ReviewDecision.ADMITTED, supersedes=None):
    return ReviewedDecision(name, SID, "reviewer", NOW, decision, "synthetic approval",
                            supersedes=supersedes)


def test_disconnected_cycle_is_not_a_resolved_review_history():
    result = decide_with([review("a", supersedes="b"), review("b", supersedes="a"), review("c")])
    assert result.decision is AdmissionDecision.REVIEW_REQUIRED


def test_unknown_review_enum_cannot_be_treated_as_approval():
    try:
        result = decide_with([review("pending-review", decision="pending")])
    except ValueError:
        return  # Rejecting malformed input at construction is also acceptable.
    assert result.decision is not AdmissionDecision.IN_SCOPE


def test_pdf_units_preserve_interleaved_heading_body_order(tmp_path):
    path = tmp_path / "sections.pdf"
    with pymupdf.open() as doc:
        page = doc.new_page()
        for y, text, size in [(72, "Section A", 18), (110, "A body 1", 11),
                              (145, "A body 2", 11), (215, "Section B", 18),
                              (250, "B body 1", 11), (285, "B body 2", 11)]:
            page.insert_text((72, y), text, fontsize=size)
        doc.save(str(path))
    result = read_document(path)
    actual = [u.raw_text for u in result.units]
    assert actual.index("A body 1") < actual.index("Section B"), actual


def test_pdf_image_region_is_accounted_for_even_with_text(tmp_path):
    path = tmp_path / "mixed-page.pdf"
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_text((72, 72), "Report title", fontsize=11)
        page.insert_image(pymupdf.Rect(72, 120, 520, 700),
                          pixmap=pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 10, 10)))
        doc.save(str(path))
    result = read_document(path)
    assert result.issues or any(u.status.value != "kept" for u in result.units), result


def test_docx_nested_table_is_extracted_or_recorded_as_gap(tmp_path):
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Outer cell"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Revenue 1234"
    path = tmp_path / "nested.docx"
    doc.save(str(path))
    result = read_document(path)
    assert "Revenue 1234" in "\n".join(u.raw_text for u in result.units) or result.issues, result


@pytest.mark.parametrize("clean,mapping", [("", ()), ("21", ((0, 1), (1, 0)))])
def test_clean_verifier_rejects_erasure_and_reordering(clean, mapping):
    with pytest.raises(CleanError):
        verify_clean_region("12", clean, mapping)


def test_heading_only_document_is_preserved(tmp_path):
    path = tmp_path / "heading.md"
    path.write_text("# Heading only\n", encoding="utf-8")
    reader = read_document(path)
    result = chunk_clean_result(reader, clean_reader_result(reader))
    assert len(result.chunks) == 1
    assert "Heading only" in result.chunks[0].search_text


def test_adjacent_markdown_tables_are_not_fused(tmp_path):
    path = tmp_path / "tables.md"
    path.write_text("| 2025 | revenue |\n|---|---|\n| A | 100 |\n\n"
                    "| 2026 | margin |\n|---|---|\n| B | 20% |\n", encoding="utf-8")
    reader = read_document(path)
    result = chunk_clean_result(reader, clean_reader_result(reader))
    tables = [chunk for chunk in result.chunks if chunk.kind == "table"]
    assert len(tables) == 2, tables


def test_table_header_context_does_not_make_valid_rows_unprocessable(tmp_path):
    path = tmp_path / "long-rows.md"
    path.write_text("| " + "H" * 990 + " |\n|---|\n| " + "A" * 990 + " |\n", encoding="utf-8")
    reader = read_document(path)
    result = chunk_clean_result(reader, clean_reader_result(reader))
    assert result.chunks
    assert all(len(c.search_text) <= 1800 or c.review_reasons for c in result.chunks)


def publish_via_job(store, build, decision_id="d0", activated_at=NOW):
    """C1 fail-closed ownership gate: publish must go through the engine job protocol.

    Protocol upgrade only; assertions unchanged. Acquire uses the real wall clock
    because MemoryStore's default clock judges lease expiry against real now.
    """
    lease = LeaseConfig(120, 30, 360, 3)
    store.register_job(build.build_id, JobStage.PUBLISHED)
    acquired = store.acquire_job(build.build_id, JobStage.PUBLISHED, "fixture",
                                 datetime.now(UTC), lease)
    token = acquired.fence_token
    assert token is not None
    return store.publish(SID, decision_id, build.build_id, activated_at,
                         owner_id="fixture", fence_token=token)


def test_control_valid_admission_and_current_publish():
    assert decide_with([review("valid")]).decision is AdmissionDecision.IN_SCOPE
    store, _, build = setup_store()
    assert publish_via_job(store, build).active_build_id == build.build_id


def test_control_connected_review_chain_resolves():
    result = decide_with([review("a"), review("b", supersedes="a"),
                          review("c", ReviewDecision.EXCLUDED, supersedes="b")])
    assert result.decision is AdmissionDecision.EXCLUDED_BY_POLICY


def test_control_single_heading_pdf_order(tmp_path):
    path = tmp_path / "single.pdf"
    with pymupdf.open() as doc:
        page = doc.new_page()
        for y, text, size in [(72, "Section A", 18), (110, "A body 1", 11),
                              (145, "A body 2", 11)]:
            page.insert_text((72, y), text, fontsize=size)
        doc.save(str(path))
    assert [u.raw_text for u in read_document(path).units] == ["Section A", "A body 1", "A body 2"]


def test_control_short_table_header(tmp_path):
    path = tmp_path / "short-header.md"
    path.write_text("| H |\n|---|\n| " + "A" * 990 + " |\n", encoding="utf-8")
    reader = read_document(path)
    result = chunk_clean_result(reader, clean_reader_result(reader))
    assert result.chunks and all(len(c.search_text) <= 1800 for c in result.chunks)


def test_control_clean_identity_mapping():
    verify_clean_region("12", "12", ((0, 0),))
