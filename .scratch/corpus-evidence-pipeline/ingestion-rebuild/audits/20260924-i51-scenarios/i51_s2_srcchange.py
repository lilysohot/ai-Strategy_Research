"""I5-1 场景二：源变化（全 8 活动源的变异副本，隔离库 i51_s2_srcchg）。

U 2026-09-24 批准范围：源变化场景在隔离环境执行，不触碰生产活动来源、不改 data/
任何字节（守卫 protected_roots=data）。变异副本落审计工作区 s2-mutated/：

- 6 份 PDF：pymupdf 元数据变更（producer 追加标记）——字节变化（source_id 变）而
  正文不变（解析产物应全等：版本翻新、内容稳定的对照面）；
- MD：追加一行正文（内容真实变化）；
- DOCX：追加一个段落（内容真实变化）。

验证链：
1. 基线：同库先复现 v1（同源同版本，build_id 与生产 I4-4 全等）并发布 8/8；
2. 版本：变异副本 → 新 source_id（≠ 原 source_id）→ 新 build_id（≠ v1）；
   6 份 PDF 的 units/chunks 与 v1 全等（元数据级变化）；MD/DOCX 未及 build——
   **dev lane 材料放宽按 source_id 白名单生效（admission.py dev_lane_sources），
   变异副本 source_id 不在白名单 ⇒ 准入 fail-closed 拒绝（expected）**，旧版本不受影响；
3. 缓存：变异副本的解析检查点按新 source_id 键控 ⇒ reparse_reason=checkpoint_absent
   （跨源版本零缓存污染）；
4. 失败恢复：4 份 gap 源的变异 build 无已签署 gap-review ⇒ check 阻断、publish 拒绝
   （fail-closed），原 source 活动版本不受影响；
5. 旧版本保留：v1 build 与 8 份活动发布在变异构建后逐源原样；
6. 生产只读前后快照全等。

产物 write-once：i51-s2-srcchange-report.json。
"""
from __future__ import annotations

import hashlib
import json
import shutil
import sys
from datetime import UTC, datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import i51_common as common  # noqa: E402

OUT = HERE / "i51-s2-srcchange-report.json"
DB = "i51_s2_srcchg"
OWNER = "i51-s2-srcchange"
NOW = datetime(2026, 9, 24, 3, 0, tzinfo=UTC)
MUTATED = HERE / "s2-mutated"
S2_ARCHIVE = HERE / "s2-archive"


def mutate_copy(orig: Path, dest: Path, fmt: str) -> dict:
    """生成变异副本（不改 data/ 任何字节），返回变异说明与前后哈希。"""
    before = hashlib.sha256(orig.read_bytes()).hexdigest()
    if fmt == "pdf":
        import pymupdf  # noqa: PLC0415

        doc = pymupdf.open(orig)
        meta = dict(doc.metadata or {})
        meta["producer"] = f"{meta.get('producer') or ''}[i5s1-scenario2-mutated]".strip()
        doc.set_metadata(meta)
        doc.save(dest)
        doc.close()
        method = "pdf_metadata_producer_append（字节变化，正文不变）"
    elif fmt == "md":
        text = orig.read_text(encoding="utf-8")
        text += "\n\n## I5-1 场景二变异附录\n\n来源变化验证用段落（2026-09-24）：本文为 I5-1 场景二变异副本，用于验证源变化后的版本/缓存/失败恢复行为。\n"
        dest.write_text(text, encoding="utf-8")
        method = "md_append_paragraph（正文真实变化）"
    elif fmt == "docx":
        import docx  # noqa: PLC0415

        d = docx.Document(str(orig))
        d.add_paragraph("I5-1 场景二变异段：来源变化验证附录（2026-09-24）。")
        d.save(str(dest))
        method = "docx_add_paragraph（正文真实变化）"
    else:  # pragma: no cover
        raise SystemExit(f"未知格式 {fmt}")
    after = hashlib.sha256(dest.read_bytes()).hexdigest()
    return {"method": method, "sha256_before": before, "sha256_after": after,
            "bytes_changed": before != after}


def main() -> int:
    if OUT.exists():
        raise SystemExit(f"拒绝：write-once 报告已存在 {OUT}")

    creds = common.load_env_credentials()
    pdsn = common.prod_dsn(creds)
    prod_pre = common.prod_readonly_snapshot(pdsn)

    cli, resolved = common.install_guard(DB)
    idsn = common.iso_dsn(creds, DB)

    manifest = common.MANIFEST
    dev_policy = json.loads(common.DEV_POLICY.read_text(encoding="utf-8"))
    dev_by_path = {str(item["path"]): item for item in dev_policy["dev_lane"]["sources"]}
    sources = json.loads(manifest.read_text(encoding="utf-8"))["sources"]
    i44 = json.loads(common.I44_REPORT.read_text(encoding="utf-8"))
    baseline_rows = {r["original_name"]: r for r in i44["per_source"]}

    from plugins.corpus.preparation.contract import sha256_of_bytes  # noqa: PLC0415
    from plugins.corpus.preparation.repository_pg import PgStore  # noqa: PLC0415

    record: dict = {
        "artifact": "i51-s2-srcchange",
        "scenario": "源变化（8 份变异副本；隔离库；不改 data/ 字节）",
        "authority": {"guard_sha256": common.digest(common.GUARD),
                      "manifest_sha256": common.digest(manifest),
                      "scope": "U 2026-09-24 批准：全四场景×全部 8 个活动源（r5n m8_kickoff）"},
        "target": {"dsn": f"postgresql://***@{common.ISO_HOST}:{common.ISO_PORT}/{DB}",
                   "resolve_target_db": resolved},
        "prod_pre": prod_pre,
    }

    # ── Step 1: v1 基线（同源同版本 → 与生产对照锚点）──
    store = PgStore(idsn, sandbox_db=resolved)
    try:
        record["seeded_decisions_v1"] = common.seed_decisions(
            store, sources, dev_by_path, creds, decision_prefix=None, now=NOW,
            rationale_tag="I5-1 场景二 v1 基线重播种")
    finally:
        store.close()
    v1, v1_rows = common.build_check_publish(
        cli, idsn, manifest, S2_ARCHIVE / "v1", OWNER, baseline_rows, gap_review=True)
    record["v1_run"] = {"build": v1["build"],
                        "built": sum(1 for r in v1_rows if r.get("build_id")),
                        "build_id_match_i44": all(r["i44_baseline"]["build_id_match"]
                                                  for r in v1_rows if r.get("build_id")),
                        "published": sum(1 for r in v1_rows
                                         if (r.get("publish") or {}).get("exit_code") == 0)}
    v1_by_name = {r["original_name"]: r for r in v1_rows}

    # ── Step 2: 生成 8 份变异副本（工作区内；data/ 零触碰）──
    if MUTATED.exists():
        shutil.rmtree(MUTATED)
    MUTATED.mkdir(parents=True)
    mutations: dict[str, dict] = {}
    mutated_paths: dict[str, str] = {}
    for entry in sources:
        orig_rel = str(entry["path"])
        orig = common.ROOT / orig_rel
        dest = MUTATED / orig.name
        fmt = orig.suffix.lower().lstrip(".")
        mutations[orig_rel] = mutate_copy(orig, dest, fmt)
        mutated_paths[orig_rel] = str(dest.relative_to(common.ROOT))
    record["mutations"] = mutations
    record["data_dir_untouched"] = True  # 守卫 protected_roots=data 强制；副本全部落 s2-mutated/

    # ── Step 3: 播种变异副本决定链 + 写 s2 manifest（新 id 同步）──
    store = PgStore(idsn, sandbox_db=resolved)
    try:
        seeded = common.seed_decisions(
            store, sources, dev_by_path, creds, decision_prefix="i5s2", now=NOW,
            rationale_tag="I5-1 场景二变异副本决定链（属性逐字保留，绑定新 source_id）",
            path_override=mutated_paths)
    finally:
        store.close()
    s2_sources = []
    for item in seeded:
        orig_entry = next(e for e in sources if str(e["path"]) == item["orig_path"])
        s2_sources.append({
            "path": item["path"],
            "domain_hint": orig_entry["domain_hint"],
            "review_decision_ids": item["decisions"],
            "provenance": f"i5s1-scenario2-mutated-fixture-of:{item['orig_path']}",
        })
    s2_manifest = HERE / "s2-manifest.json"
    common.write_manifest(
        s2_manifest, s2_sources,
        note="I5-1 场景二：8 份变异副本（provenance 字段逐源回指原件；决定 id 为 i5s2 前缀新链）")
    record["s2_manifest_sha256"] = common.digest(s2_manifest)

    # ── Step 4: 构建变异副本 → 版本/准入/缓存行为 ──
    v2, v2_rows = common.build_check_publish(
        cli, idsn, s2_manifest, S2_ARCHIVE / "v2", OWNER, {}, gap_review=False)
    record["v2_build"] = v2["build"]
    rows = []
    for row in v2_rows:
        v1r = v1_by_name[row["original_name"]]
        built = row.get("decision") == "in_scope" and row.get("build_id")
        item = {
            "original_name": row["original_name"],
            "decision": row.get("decision"),
            "new_source_id_differs": bool(row.get("source_id")
                                          and row["source_id"] != v1r.get("source_id")),
            "new_build_id_differs": bool(built and row.get("build_id") != v1r.get("build_id")),
            "units_equal_v1": (row.get("unit_count") == v1r.get("unit_count")) if built else None,
            "chunks_equal_v1": (row.get("chunk_count") == v1r.get("chunk_count")) if built else None,
            "publish_exit": (row.get("publish") or {}).get("exit_code"),
            "publish_blocked": bool((row.get("publish") or {}).get("blocked")),
        }
        rows.append(item)
    record["v2_outcomes"] = rows

    # ── Step 5: 缓存键控核对（变异副本检查点 = checkpoint_absent；新 source_id 命名空间）──
    store = PgStore(idsn, sandbox_db=resolved)
    try:
        ckpt_rows = []
        for row in v2_rows:
            if row.get("source_id"):
                ck = store.get_source_checkpoint(row["source_id"], "parse")
                if ck is not None:
                    data = json.loads(ck)
                    ckpt_rows.append({"original_name": row["original_name"],
                                      "reparse_reason": data.get("reparse_reason"),
                                      "keyed_by_mutated_source_id": True})
        record["v2_checkpoints"] = ckpt_rows

        # ── Step 6: 旧版本保留核对（v1 build 与活动发布逐原样）──
        old_retained = []
        for name, v1r in v1_by_name.items():
            if not v1r.get("build_id"):
                continue
            b = store.get_build(v1r["build_id"])
            pub = store.get_publication(b.source_id) if b else None
            old_retained.append({
                "original_name": name,
                "v1_build_present": b is not None,
                "v1_still_active": bool(pub and pub.active_build_id == v1r["build_id"]),
                "generation_unchanged": (pub.generation == v1r.get("generation")) if pub else None,
            })
        record["v1_retained"] = old_retained
    finally:
        store.close()
    record["iso_counts_final"] = common.iso_counts(idsn)

    built_v2 = [r for r in rows if r["decision"] == "in_scope"]
    excluded_v2 = [r for r in rows if r["decision"] != "in_scope"]
    record["summary"] = {
        "v1_built": record["v1_run"]["built"],
        "v1_build_id_match_i44": record["v1_run"]["build_id_match_i44"],
        "v1_published": record["v1_run"]["published"],
        "v2_in_scope": len(built_v2),
        "v2_excluded_fail_closed": len(excluded_v2),
        "v2_new_source_ids": all(r["new_source_id_differs"] for r in rows if r["decision"] == "in_scope"),
        "v2_new_build_ids": all(r["new_build_id_differs"] for r in built_v2),
        "v2_units_chunks_equal_v1": all(r["units_equal_v1"] and r["chunks_equal_v1"]
                                        for r in built_v2),
        "v2_gap_blocked": sum(1 for r in built_v2 if r["publish_blocked"]),
        "v2_published_new": sum(1 for r in built_v2 if r["publish_exit"] == 0),
        "v2_checkpoints_fresh": all(c["reparse_reason"] == "checkpoint_absent"
                                    for c in record["v2_checkpoints"]),
        "v1_retained_active": all(r["v1_build_present"] and r["v1_still_active"]
                                  for r in record["v1_retained"]),
        "mutations_all_bytes_changed": all(m["bytes_changed"] for m in mutations.values()),
    }
    record["prod_post"] = common.prod_readonly_snapshot(pdsn)
    record["prod_unchanged"] = record["prod_post"] == prod_pre

    gate = (
        record["summary"]["v1_built"] == 8 and record["summary"]["v1_build_id_match_i44"]
        and record["summary"]["v1_published"] == 8
        and record["summary"]["v2_in_scope"] == 6
        and record["summary"]["v2_excluded_fail_closed"] == 2
        and record["summary"]["v2_new_source_ids"] and record["summary"]["v2_new_build_ids"]
        and record["summary"]["v2_units_chunks_equal_v1"]
        and record["summary"]["v2_gap_blocked"] == 4 and record["summary"]["v2_published_new"] == 2
        and record["summary"]["v2_checkpoints_fresh"]
        and record["summary"]["v1_retained_active"]
        and record["summary"]["mutations_all_bytes_changed"]
        and record["prod_unchanged"]
    )
    record["gate_passed"] = gate
    record["gate_note"] = ("通过：6 源变化 → 新 source/build（元数据级变化、产物全等）；"
                           "2 dev lane 变异副本按 source_id 白名单 fail-closed 拒绝（设计内）；"
                           "4 gap 源变异 build 发布阻断且旧版本照常服务；v1 全保留；生产零触碰"
                           if gate else "未通过")

    OUT.write_text(json.dumps(record, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"gate_passed": gate, "summary": record["summary"]},
                     ensure_ascii=False, indent=2))
    print("sha256:", hashlib.sha256(OUT.read_bytes()).hexdigest()[:16], "…")
    return 0 if gate else 1


if __name__ == "__main__":
    raise SystemExit(main())
