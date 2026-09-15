"""Versioned evidence packets: immutable source text plus verifiable layout coordinates.

The page remains the archival unit. Extraction consumes complete paragraphs or
table rows with repeated headers; a low-text image page is explicitly unknown.
"""

from __future__ import annotations

import hashlib
import json
import re
from itertools import pairwise
from pathlib import Path
from typing import cast

import pymupdf
from docx import Document
from docx.table import Table
from pydantic import BaseModel, ConfigDict

from plugins.corpus.claims import document_ticker
from plugins.corpus.ingest import _doc_id, _title_from_filename, content_hash
from plugins.corpus.metadata import derive_published

PARSER_VERSION = "evidence-layout-4"
PERIOD = re.compile(r"20\d{2}(?:[AEF]|H[12]|Q[1-4])?", re.IGNORECASE)
NUMBER = re.compile(r"[+\-−]?(?:\d[\d,]*(?:\.\d+)?|\(\d[\d,]*(?:\.\d+)?\))%?")
_QUESTION_TURN_BOUNDARY = re.compile(
    r"(?m)^(?=\s*(?:主持人|投资者|提问者|分析师|问)\s*[：:])"
)
_SPEAKER_TURN_BOUNDARY = re.compile(
    r"(?m)^(?=\s*(?:主持人|专家|投资者|提问者|回答者|管理层|分析师|嘉宾)\s*[：:])"
)
_NUMBERED_QUESTION_BOUNDARY = re.compile(
    r"(?m)^(?=\s*(?:[一二三四五六七八九十百]+|\d+)[、.．]\s*[^\n]{0,100}[？?])"
)
BBox = tuple[float, float, float, float]


def fingerprint(value: object) -> str:
    """Content identity independent of file location and timestamps."""
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=False, sort_keys=True, default=str).encode()
    ).hexdigest()


class Span(BaseModel):
    model_config = ConfigDict(frozen=True)
    locator: str
    text: str
    bbox: BBox | None = None
    start: int | None = None
    end: int | None = None


class Cell(BaseModel):
    model_config = ConfigDict(frozen=True)
    table_id: str
    row: str
    column: str
    value: str
    unit: str
    span: Span
    row_span: Span
    column_span: Span
    unit_span: Span | None = None

    def reference(self) -> dict[str, str]:
        return {"table": self.table_id, "row": self.row, "column": self.column, "cell": self.value}


class EvidencePacket(BaseModel):
    model_config = ConfigDict(frozen=True)
    packet_id: str
    locator: str
    kind: str
    text: str
    spans: tuple[Span, ...] = ()
    cells: tuple[Cell, ...] = ()
    context: tuple[str, ...] = ()
    status: str = "available"
    reasons: tuple[str, ...] = ()

    def lookup(self, reference: dict[str, str]) -> str | None:
        """Resolve only a complete, exact reference to a parsed source cell."""
        for cell in self.cells:
            if reference == cell.reference():
                return cell.value
        return None


class EvidenceDocument(BaseModel):
    model_config = ConfigDict(frozen=True)
    doc_id: str
    title: str
    source_path: str
    source_rev: str
    parse_rev: str
    parser_version: str
    subject: str | None
    published: str | None
    pages: tuple[Span, ...]
    packets: tuple[EvidencePacket, ...]

    def fetch(self, packet_id: str) -> EvidencePacket:
        return next(p for p in self.packets if p.packet_id == packet_id)


def _packet(parse_rev: str, **fields: object) -> EvidencePacket:
    payload = {k: v for k, v in fields.items()}
    # Pydantic JSON encoders give spans/cells a stable structural representation.
    preliminary = EvidencePacket(packet_id="", **payload)  # type: ignore[arg-type]
    key = fingerprint([parse_rev, preliminary.model_dump(mode="json")])
    return preliminary.model_copy(update={"packet_id": key})


def split_spans(text: str, locator: str, max_chars: int) -> list[Span]:
    """Cover every character once without separating ordinary question/answer turns."""
    if max_chars < 100:
        raise ValueError("packet_chars must be at least 100")
    spans: list[Span] = []
    start = 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        if end < len(text):
            lower = start + max_chars // 2
            structured: list[int] = []
            for pattern in (
                _QUESTION_TURN_BOUNDARY,
                _NUMBERED_QUESTION_BOUNDARY,
                _SPEAKER_TURN_BOUNDARY,
            ):
                structured = [match.start() for match in pattern.finditer(text, lower, end)]
                if structured:
                    break
            # Prefer the latest question-start boundary so its answer stays in the next packet.
            # If no structural marker exists, retain the established lossless text fallback.
            boundary = max(structured, default=-1)
            if boundary < 0:
                choices = [
                    text.rfind(sep, lower, end) for sep in ("\n\n", "\n", "。", ". ")
                ]
                boundary = max(choices)
            if boundary >= 0:
                end = boundary if structured else boundary + 1
        spans.append(Span(locator=locator, text=text[start:end], start=start, end=end))
        start = end
    return spans


def _lines(page: pymupdf.Page) -> list[list[Span]]:
    words = sorted(
        cast(list[tuple[float, float, float, float, str, int, int, int]], page.get_text("words")),
        key=lambda w: (w[1], w[0]),
    )
    assert page.number is not None
    lines: list[list[Span]] = []
    for word in words:
        span = Span(
            locator=str(page.number + 1), text=word[4], bbox=(word[0], word[1], word[2], word[3])
        )
        if not lines or abs(word[1] - lines[-1][0].bbox[1]) > 2.5:  # type: ignore[index]
            lines.append([span])
        else:
            lines[-1].append(span)
    return [sorted(line, key=lambda s: s.bbox[0]) for line in lines]  # type: ignore[index]


def _join(spans: list[Span]) -> Span:
    boxes = [s.bbox for s in spans if s.bbox is not None]
    return Span(
        locator=spans[0].locator,
        text=" ".join(s.text for s in spans),
        bbox=(
            min(b[0] for b in boxes),
            min(b[1] for b in boxes),
            max(b[2] for b in boxes),
            max(b[3] for b in boxes),
        ),
    )


def _row_unit(label: str, values: list[Span], header_unit: str) -> str:
    if all(v.text.endswith("%") for v in values):
        return "%"
    if label.upper().replace(" ", "") in {"P/E", "PE", "P/B", "PB", "EV/EBITDA"}:
        return "倍"
    if any(s in label for s in ("市盈率", "市净率")):
        return "倍"
    inline = re.search(r"[（(](百万元|亿元|万元|元)[）)]", label)
    if inline and not any(s in label for s in ("每股", "EPS")):
        return inline.group(1)
    if "天数" in label:
        return "天"
    if any(s in label for s in ("每股", "EPS")) and ("元" in label or header_unit):
        return "元/股"
    if label in {"流动比率", "速动比率", "总资产周转率"}:
        return "倍"
    return header_unit


def _table_packets(page: pymupdf.Page, parse_rev: str) -> list[EvidencePacket]:
    """Recognize monotonic period headers and geometrically aligned numeric rows.

    Conservative adapter for financial forecast tables, including borderless
    two-column pages. Unsupported tables remain available as prose, never guessed.
    """
    lines = _lines(page)
    assert page.number is not None
    packets: list[EvidencePacket] = []
    for index, line in enumerate(lines):
        periods = [s for s in line if PERIOD.fullmatch(s.text)]
        runs: list[list[Span]] = []
        for span in periods:
            if not runs or span.text[:4] <= runs[-1][-1].text[:4]:
                runs.append([span])
            else:
                runs[-1].append(span)
        for headers in runs:
            if len(headers) < 2:
                continue
            centers = [(h.bbox[0] + h.bbox[2]) / 2 for h in headers]  # type: ignore[index]
            spacing = min(b - a for a, b in pairwise(centers))
            if spacing < 10:
                continue
            before = [s for s in line if s.bbox[2] < centers[0] - spacing / 2]  # type: ignore[index]
            previous_run = [s for s in periods if s.bbox[0] < headers[0].bbox[0]]  # type: ignore[index]
            floor = previous_run[-1].bbox[2] if previous_run else 0  # type: ignore[index]
            labels = [s for s in before if s.bbox[0] >= floor and not PERIOD.fullmatch(s.text)]  # type: ignore[index,operator]
            left = labels[-1].bbox[0] if labels else max(floor, centers[0] - spacing * 2.5)  # type: ignore[index, type-var]
            right = centers[-1] + spacing / 2
            upper = headers[0].bbox[1]  # type: ignore[index]
            nearby = [
                s
                for prev in lines[:index]
                for s in prev
                if s.bbox is not None
                and left <= s.bbox[0] < centers[0]
                and 0 < upper - s.bbox[1] < 32
            ]  # type: ignore[index,operator]
            title = " ".join(s.text for s in nearby if not NUMBER.fullmatch(s.text))
            unit_span = next((s for s in labels if "单位" in s.text), None)
            unit_match = re.search(
                r"单位\s*[:：]\s*(百万元|亿元|万元|元|亿美元|美元)",
                unit_span.text if unit_span else "",
            )
            unit = unit_match.group(1) if unit_match else ""
            table_id = fingerprint([parse_rev, page.number, [h.model_dump() for h in headers]])
            for row in lines[index + 1 :]:
                region = [s for s in row if left <= s.bbox[0] < right]  # type: ignore[index,operator]
                if not region:
                    continue
                if sum(bool(PERIOD.fullmatch(s.text)) for s in region) >= 2:
                    break
                row_labels = [
                    s
                    for s in region
                    if s.bbox is not None
                    and s.bbox[2] < centers[0] - spacing / 2
                    and not NUMBER.fullmatch(s.text)
                ]  # type: ignore[index]
                if not row_labels or any("资料来源" in s.text for s in row_labels):
                    continue
                values: list[Span] = []
                for center in centers:
                    matches = [
                        s
                        for s in region
                        if s.bbox is not None
                        and NUMBER.fullmatch(s.text)
                        and abs((s.bbox[0] + s.bbox[2]) / 2 - center) < spacing / 2
                    ]  # type: ignore[index]
                    if len(matches) != 1:
                        break
                    values.append(matches[0])
                if len(values) != len(headers):
                    continue
                label = _join(row_labels)
                row_unit = _row_unit(label.text, values, unit)
                cells = tuple(
                    Cell(
                        table_id=table_id,
                        row=label.text,
                        column=h.text,
                        value=v.text,
                        unit=row_unit,
                        span=v,
                        row_span=label,
                        column_span=h,
                        unit_span=unit_span,
                    )
                    for h, v in zip(headers, values, strict=True)
                )
                text = " | ".join([label.text, *[f"{c.column}: {c.value}" for c in cells]])
                packets.append(
                    _packet(
                        parse_rev,
                        locator=str(page.number + 1),
                        kind="table",
                        text=text,
                        cells=cells,
                        spans=tuple([label, *values]),
                        context=tuple(x for x in (title, unit_span.text if unit_span else "") if x),
                    )
                )
    return packets


def _customer_packets(page: pymupdf.Page, parse_rev: str) -> list[EvidencePacket]:
    """Recover customer sales/share columns, with title and footer evidence.

    Explicit header and per-row unit checks are required. No OCR, interpolation,
    missing-cell filling, or inference that a customer amount is company revenue.
    """
    lines = _lines(page)
    packets: list[EvidencePacket] = []
    for index, line in enumerate(lines):
        label_header = next((s for s in line if s.text == "客户名称"), None)
        sales = next(
            (s for s in line if re.fullmatch(r"销售额[（(](百万元|亿元|万元|元)[）)]", s.text)),
            None,
        )
        share = next((s for s in line if s.text == "占收入比例"), None)
        if not label_header or not sales or not share:
            continue
        headers = (label_header, sales, share)
        if any(h.bbox is None for h in headers):
            continue
        assert sales.bbox is not None
        centers = [(h.bbox[0] + h.bbox[2]) / 2 for h in headers]  # type: ignore[index]
        if not centers[0] < centers[1] < centers[2]:
            continue
        title_parts = [
            s
            for prev in lines[:index]
            for s in prev
            if s.bbox and 0 < sales.bbox[1] - s.bbox[1] < 32
        ]  # type: ignore[index]
        title = _join(title_parts) if title_parts else None
        if title is None or "客户" not in title.text:
            continue
        unit_match = re.search(r"[（(](百万元|亿元|万元|元)[）)]", sales.text)
        assert unit_match is not None
        boundaries = ((centers[0] + centers[1]) / 2, (centers[1] + centers[2]) / 2)
        rows: list[tuple[Span, Span, Span]] = []
        footer = None
        previous_y = sales.bbox[1]  # type: ignore[index]
        for row in lines[index + 1 :]:
            joined = " ".join(s.text for s in row)
            if re.search(r"(?:数据|资料)来源", joined):
                footer = _join(row)
                break
            if not row[0].bbox or row[0].bbox[1] - previous_y > 45:
                break
            labels = [
                s
                for s in row
                if s.bbox and s.bbox[0] < boundaries[0] and not NUMBER.fullmatch(s.text)
            ]
            amounts = [
                s
                for s in row
                if s.bbox
                and boundaries[0] <= (s.bbox[0] + s.bbox[2]) / 2 < boundaries[1]
                and NUMBER.fullmatch(s.text)
                and not s.text.endswith("%")
            ]
            shares = [
                s
                for s in row
                if s.bbox
                and boundaries[1]
                <= (s.bbox[0] + s.bbox[2]) / 2
                < centers[2] + (centers[2] - centers[1]) / 2
                and NUMBER.fullmatch(s.text)
                and s.text.endswith("%")
            ]
            if not labels or len(amounts) != 1 or len(shares) != 1:
                break
            rows.append((_join(labels), amounts[0], shares[0]))
            previous_y = row[0].bbox[1]
        table_id = fingerprint(
            [parse_rev, "customer-sales", title.model_dump(), [h.model_dump() for h in headers]]
        )
        for label, amount, percent in rows:
            cells = tuple(
                Cell(
                    table_id=table_id,
                    row=label.text,
                    column=header.text,
                    value=value.text,
                    unit=unit,
                    span=value,
                    row_span=label,
                    column_span=header,
                    unit_span=value if unit == "%" else header,
                )
                for header, value, unit in ((sales, amount, unit_match[1]), (share, percent, "%"))
            )
            context = (title.text, *((footer.text,) if footer else ()))
            packets.append(
                _packet(
                    parse_rev,
                    locator=label.locator,
                    kind="table",
                    text=" | ".join(
                        [title.text, label.text, *[f"{c.column}: {c.value}" for c in cells]]
                    ),
                    cells=cells,
                    spans=(*headers, title, label, amount, percent, *((footer,) if footer else ())),
                    context=context,
                )
            )
    return packets


def parse_evidence(
    path: str | Path,
    *,
    packet_chars: int = 2000,
    pages: tuple[int, ...] | None = None,
) -> EvidenceDocument:
    """Preserve extracted text and coordinates; unsupported visual content stays unknown."""
    source = Path(path)
    digest = content_hash(source)
    title = _title_from_filename(source)
    doc_id = _doc_id(source, digest)
    parse_rev = fingerprint([digest, PARSER_VERSION, pymupdf.VersionBind, packet_chars, pages])
    raw_pages: list[Span] = []
    packets: list[EvidencePacket] = []
    if source.suffix.lower() == ".pdf":
        with pymupdf.open(source) as pdf:
            selected = list(pages) if pages is not None else list(range(1, len(pdf) + 1))
            if len(set(selected)) != len(selected) or any(n < 1 or n > len(pdf) for n in selected):
                raise ValueError("invalid page selection")
            first_text = cast(str, pdf[0].get_text("text")) if len(pdf) else ""
            for number in selected:
                page = pdf[number - 1]
                text = cast(str, page.get_text("text"))
                raw_pages.append(Span(locator=str(number), text=text))
                if len(text.strip()) < 50:
                    packets.append(
                        _packet(
                            parse_rev,
                            locator=str(number),
                            kind="visual",
                            text=text,
                            status="unknown",
                            reasons=("needs_ocr",),
                        )
                    )
                    continue
                packets.extend(_table_packets(page, parse_rev))
                packets.extend(_customer_packets(page, parse_rev))
                for span in split_spans(text, str(number), packet_chars):
                    packets.append(
                        _packet(
                            parse_rev,
                            locator=str(number),
                            kind="prose",
                            text=span.text,
                            spans=(span,),
                            context=(title,),
                        )
                    )
    else:
        if pages is not None:
            raise ValueError("page selection is only supported for PDFs")
        if source.suffix.lower() == ".md":
            # Keep line endings and offsets, including content the legacy cleaner suppresses.
            with source.open(encoding="utf-8", newline="") as stream:
                text = stream.read()
            raw_pages.append(Span(locator="document", text=text))
            for span in split_spans(text, "document", packet_chars):
                headings: dict[int, str] = {}
                for match in re.finditer(r"^(#{1,6})\s+(.+)$", text[: span.end], re.MULTILINE):
                    level = len(match.group(1))
                    headings = {k: v for k, v in headings.items() if k < level}
                    headings[level] = match.group(2).strip()
                packets.append(
                    _packet(
                        parse_rev,
                        locator="document",
                        kind="prose",
                        text=span.text,
                        spans=(span,),
                        context=(title, *headings.values()),
                    )
                )
        elif source.suffix.lower() == ".docx":
            headings = {}
            for index, item in enumerate(Document(str(source)).iter_inner_content()):
                locator = f"element:{index}"
                if isinstance(item, Table):
                    text = "\n".join(
                        " | ".join(cell.text for cell in row.cells) for row in item.rows
                    )
                else:
                    text = item.text
                    style = item.style.name if item.style else ""
                    match = re.fullmatch(r"Heading (\d)", style or "", re.IGNORECASE)
                    if match:
                        level = int(match.group(1))
                        headings = {k: v for k, v in headings.items() if k < level}
                        headings[level] = text
                raw_pages.append(Span(locator=locator, text=text))
                for span in split_spans(text, locator, packet_chars):
                    packets.append(
                        _packet(
                            parse_rev,
                            locator=locator,
                            kind="prose",
                            text=span.text,
                            spans=(span,),
                            context=(title, *headings.values()),
                        )
                    )
        else:
            raise ValueError("unsupported source format")
        first_text = "\n".join(p.text for p in raw_pages)[:4000]
        if not packets:
            packets.append(
                _packet(
                    parse_rev,
                    locator="document",
                    kind="visual",
                    text="",
                    status="unknown",
                    reasons=("no_extractable_text",),
                )
            )
    published = derive_published(doc_id, title, first_text)
    return EvidenceDocument(
        doc_id=doc_id,
        title=title,
        source_path=str(source),
        source_rev=digest,
        parse_rev=parse_rev,
        parser_version=PARSER_VERSION,
        subject=document_ticker(title, [first_text]),
        published=published.isoformat() if published else None,
        pages=tuple(raw_pages),
        packets=tuple(packets),
    )
