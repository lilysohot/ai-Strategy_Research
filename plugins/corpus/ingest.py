"""P0b B2：语料 ingest——把 PDF / DOCX / MD 变成**可定位**的原文块。

三条设计约束决定了这里的写法：

1. **定位符（locator）而不是页码**。PDF 有天然页码，DOCX 和 MD 没有。
   而 ``strategy_lint`` 要求 ``evidence.page`` 必须是真实定位符（禁止 ``"—"``），
   所以每种格式都必须给出一个能取回原文的 locator：

   ==========  ============  ==========================
   格式        locator        取值
   ==========  ============  ==========================
   PDF         物理页码       ``"1"`` ``"2"`` …
   DOCX        段落块         ``"para12"`` 或章节标题
   MD          章节标题       ``"二、关键事实"``
   ==========  ============  ==========================

2. **幂等**。以文件字节的 content_hash 为键，同一份重复 ingest 不产生
   重复文档——「故意重复的那份只入库一次」这条验收直接靠它。

3. **纯解析与落库分离**。``parse_*`` 不碰数据库，单测可以直接喂文件；
   落库是独立的一步。否则每测一次解析规则都要搭一次 DB。

扫描页只**标记**不处理（``needs_ocr``）：P0 不接 OCR，但必须让调用方知道
这份文档的文字层是空的，而不是安静地存一堆空块。
"""

from __future__ import annotations

import hashlib
import re
import sqlite3
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument

# 低于此值（字符/页）判定为扫描页：正常研报每页上千字符，扫描件接近 0。
MIN_CHARS_PER_PAGE = 50
# 落盘路径。P0 规模用单文件 SQLite 足够，不上服务端。
DEFAULT_DB_PATH = "data/corpus/index.db"
CORPUS_ROOT = "data/corpus"

STATUS_OK = "ok"
STATUS_NEEDS_OCR = "needs_ocr"
STATUS_EMPTY = "empty"

# 券商/媒体研报的固定噪音：免责声明、页眉页脚、页码行。
# 故意只收「几乎不可能是正文」的模式——剥离过头会误删有效段落，
# 而漏掉几条噪音只是索引里多点冗余，代价不对称。
BOILERPLATE_PATTERNS: tuple[str, ...] = (
    r"^\s*免责声明\s*[:：]?",
    r"^\s*分析师声明\s*[:：]?",
    r"^\s*本报告由.*仅供.*使用\s*$",
    r"^\s*本报告不构成.*投资建议\s*$",
    r"^\s*(第\s*)?\d+\s*/\s*\d+\s*(页)?\s*$",  # "3 / 22" 页码行
    r"^\s*-?\s*\d+\s*-?\s*$",  # 孤立的页码
    r"^\s*请务必阅读.*免责声明.*$",
    r"^\s*未经.*书面许可.*不得.*$",
)
_BOILERPLATE_RE = tuple(re.compile(p) for p in BOILERPLATE_PATTERNS)


@dataclass(frozen=True)
class Block:
    """一个可定位的原文块。

    ``locator`` 是给人和模型看的定位符，``seq`` 是给数据库排序用的序号——
    两者分开是因为 DOCX/MD 的 locator 是中文标题，没法用来排序。
    """

    seq: int
    locator: str
    text: str


@dataclass
class ParsedDocument:
    """一份文档的解析结果（未落库）。"""

    doc_id: str
    title: str
    source_path: str
    content_hash: str
    mime: str
    blocks: list[Block] = field(default_factory=list)
    status: str = STATUS_OK

    @property
    def char_count(self) -> int:
        return sum(len(block.text) for block in self.blocks)


def content_hash(path: str | Path) -> str:
    """文件字节的 SHA-256（取前 16 位）。

    用字节而不是解析后的文本：同一份 PDF 换个解析器文本可能不同，
    但字节不会说谎，去重要的就是这个稳定性。
    """
    digest = hashlib.sha256()
    digest.update(Path(path).read_bytes())
    return digest.hexdigest()[:16]


def _strip_boilerplate(text: str) -> str:
    """按行剥离固定噪音，保留行结构（删行不填空行，避免破坏段落连贯）。"""
    kept = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            kept.append("")
            continue
        if any(pattern.match(stripped) for pattern in _BOILERPLATE_RE):
            continue
        kept.append(line)
    return "\n".join(kept)


def _doc_id(source_path: str | Path, file_hash: str) -> str:
    """``<日期>_<hash8>``——短、无空格、无中文，模型抄写时不容易抄错。

    为什么不用完整文件名做 id：那些名字有 60+ 个中文字符，让模型逐字复制
    到 ``evidence.source_ref`` 里，抄错一个字就成了悬空引用（ERROR）。
    source_path 里仍保留原始文件名，可回溯性没丢。
    """
    path = Path(source_path)
    match = re.match(r"(\d{4}-\d{2}-\d{2})", path.name)
    date_part = match.group(1) if match else "undated"
    return f"{date_part}_{file_hash[:8]}"


def _title_from_filename(source_path: str | Path) -> str:
    """从文件名剥出标题：去掉日期前缀、hash 后缀与扩展名。"""
    stem = Path(source_path).stem
    stem = re.sub(r"^\d{4}-\d{2}-\d{2}[_\s]*", "", stem)
    stem = re.sub(r"[-_][0-9a-f]{8}$", "", stem)
    return stem.strip(" -_") or stem


def parse_pdf(path: str | Path) -> list[Block]:
    """PDF → 每页一个块，locator 是 1-based 页码字符串。

    用 ``pymupdf``（而非 ``fitz`` 旧入口，也非 pypdf）：只有它提供版面坐标，
    且实测这 14 份都能直接抽出文字层，无需 OCR。
    """
    blocks: list[Block] = []
    with pymupdf.open(str(path)) as doc:
        for index, page in enumerate(doc, start=1):
            text = _strip_boilerplate(page.get_text("text"))
            blocks.append(Block(seq=index, locator=str(index), text=text))
    return blocks


def parse_docx(path: str | Path) -> list[Block]:
    """DOCX → 按「标题切块」。

    优先用 Word 的 Heading 样式切块（locator = 标题文字，语义最好）；
    没有标题样式时退化为按段落序号每 20 段一块——总比整篇一个块强，
    因为一个块等于整个文档的话，取证就退化成「全文搜索」，页码形同虚设。
    """
    document = DocxDocument(str(path))
    blocks: list[Block] = []
    current_locator: str | None = None
    buffer: list[str] = []
    seq = 0
    fallback_index = 0

    def flush() -> None:
        nonlocal buffer, seq
        text = _strip_boilerplate("\n".join(buffer)).strip()
        if text:
            locator = current_locator or f"para{fallback_index}"
            blocks.append(Block(seq=seq, locator=locator, text=text))
            seq += 1
        buffer = []

    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        style = (paragraph.style.name or "") if paragraph.style is not None else ""
        if style.lower().startswith("heading") and text:
            flush()
            current_locator = text
            fallback_index = index
            continue
        if text:
            buffer.append(text)
        # 无标题样式时按段落数兜底切块
        if current_locator is None and len(buffer) >= 20:
            fallback_index = index
            flush()
    flush()
    return blocks


def parse_markdown(path: str | Path) -> list[Block]:
    """MD → 按 ``#`` 标题切块，locator 是标题文字（去掉 # 号）。

    标题为空（首段在第一个标题之前）时给 ``开头``，保证 locator 非空——
    空 locator 会被 lint 判为占位符。
    """
    blocks: list[Block] = []
    current_locator = "开头"
    buffer: list[str] = []
    seq = 0

    def flush() -> None:
        nonlocal buffer, seq
        text = _strip_boilerplate("\n".join(buffer)).strip()
        if text:
            blocks.append(Block(seq=seq, locator=current_locator, text=text))
            seq += 1
        buffer = []

    for line in Path(path).read_text(encoding="utf-8").splitlines():
        if re.match(r"^#{1,6}\s+", line):
            flush()
            current_locator = line.lstrip("#").strip() or current_locator
            continue
        buffer.append(line)
    flush()
    return blocks


_PARSERS = {
    ".pdf": (parse_pdf, "application/pdf"),
    ".docx": (
        parse_docx,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ),
    ".md": (parse_markdown, "text/markdown"),
}


def parse_document(path: str | Path) -> ParsedDocument:
    """解析一份文档。纯函数：不写库、不产生副作用。"""
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix not in _PARSERS:
        raise ValueError(
            "不支持的语料格式：{}（支持 {}）".format(
                suffix,
                "、".join(sorted(_PARSERS)),
            )
        )
    parser, mime = _PARSERS[suffix]
    file_hash = content_hash(source)
    blocks = parser(source)

    total_chars = sum(len(block.text) for block in blocks)
    if total_chars == 0:
        status = STATUS_EMPTY
    elif total_chars / max(len(blocks), 1) < MIN_CHARS_PER_PAGE:
        status = STATUS_NEEDS_OCR
    else:
        status = STATUS_OK

    return ParsedDocument(
        doc_id=_doc_id(source, file_hash),
        title=_title_from_filename(source),
        source_path=str(source),
        content_hash=file_hash,
        mime=mime,
        blocks=blocks,
        status=status,
    )


def iter_corpus_files(root: str | Path = CORPUS_ROOT) -> Iterable[Path]:
    """遍历语料目录里的可解析文件（跳过 README 与隐藏文件）。"""
    directory = Path(root)
    for path in sorted(directory.iterdir()):
        if not path.is_file() or path.name.startswith("."):
            continue
        if path.name == "README.md":
            continue
        if path.suffix.lower() in _PARSERS:
            yield path


# content_hash 上加 UNIQUE：幂等由数据库保证，而不是靠调用方记得先查一次。
# 「故意重复的那份只入库一次」这条验收，本质就是这个约束在兜底。
SCHEMA = """
CREATE TABLE IF NOT EXISTS documents (
    doc_id       TEXT PRIMARY KEY,
    title        TEXT NOT NULL,
    source_path  TEXT NOT NULL,
    content_hash TEXT NOT NULL UNIQUE,
    mime         TEXT NOT NULL,
    status       TEXT NOT NULL,
    char_count   INTEGER NOT NULL DEFAULT 0,
    block_count  INTEGER NOT NULL DEFAULT 0,
    ingested_at  TEXT NOT NULL DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS blocks (
    doc_id  TEXT NOT NULL REFERENCES documents(doc_id) ON DELETE CASCADE,
    seq     INTEGER NOT NULL,
    locator TEXT NOT NULL,
    text    TEXT NOT NULL,
    PRIMARY KEY (doc_id, seq)
);

CREATE INDEX IF NOT EXISTS idx_blocks_doc ON blocks(doc_id);
"""


def connect(db_path: str | Path = DEFAULT_DB_PATH) -> sqlite3.Connection:
    """打开（必要时创建）语料库。外键约束显式打开，否则 ON DELETE CASCADE 不生效。"""
    path = Path(db_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(path))
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db(conn: sqlite3.Connection) -> None:
    """建表。可重复调用（``IF NOT EXISTS``）。"""
    conn.executescript(SCHEMA)
    conn.commit()


def upsert_document(conn: sqlite3.Connection, parsed: ParsedDocument) -> bool:
    """写库。返回 ``True`` 表示新入库，``False`` 表示 content_hash 已存在被跳过。

    幂等的落点在 ``content_hash UNIQUE`` 上：捕获 IntegrityError 而不是先 SELECT
    再 INSERT——后者在并发/重复调用之间有竞态，前者由数据库原子保证。
    """
    try:
        with conn:
            conn.execute(
                "INSERT INTO documents (doc_id, title, source_path, content_hash,"
                " mime, status, char_count, block_count)"
                " VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    parsed.doc_id,
                    parsed.title,
                    parsed.source_path,
                    parsed.content_hash,
                    parsed.mime,
                    parsed.status,
                    parsed.char_count,
                    len(parsed.blocks),
                ),
            )
    except sqlite3.IntegrityError:
        return False

    with conn:
        conn.executemany(
            "INSERT INTO blocks (doc_id, seq, locator, text) VALUES (?, ?, ?, ?)",
            [(parsed.doc_id, block.seq, block.locator, block.text) for block in parsed.blocks],
        )
    return True


def ingest_corpus(
    root: str | Path = CORPUS_ROOT,
    db_path: str | Path = DEFAULT_DB_PATH,
) -> dict[str, int]:
    """解析整个语料目录并落库，返回统计。

    统计里刻意分开 ``added`` 与 ``skipped_duplicate``：重跑 ingest 是常态
    （20 份样本也会换批次），如果重复被静默吞掉，就没法判断这次到底新进了几份。
    """
    conn = connect(db_path)
    init_db(conn)
    stats = {
        "total": 0,
        "added": 0,
        "skipped_duplicate": 0,
        "needs_ocr": 0,
        "empty": 0,
        "blocks": 0,
        "failed": 0,
    }
    try:
        for path in iter_corpus_files(root):
            stats["total"] += 1
            try:
                parsed = parse_document(path)
            except Exception:  # 单份失败不该中断整批
                stats["failed"] += 1
                continue
            if parsed.status == STATUS_NEEDS_OCR:
                stats["needs_ocr"] += 1
            elif parsed.status == STATUS_EMPTY:
                stats["empty"] += 1
            if upsert_document(conn, parsed):
                stats["added"] += 1
                stats["blocks"] += len(parsed.blocks)
            else:
                stats["skipped_duplicate"] += 1
    finally:
        conn.close()
    return stats
